import os
import shutil
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Establece márgenes internos de una celda en dxa."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Establece color de fondo de una celda."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Aplica bordes limpios y discretos a la tabla."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def create_styled_document():
    doc = Document()

    # Configuración de página: Carta con márgenes de 2.0 cm (aprox 0.79 pulgadas)
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Estilos tipográficos base
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x26, 0x26, 0x26)

    # Paleta de colores institucionales sobrios
    COLOR_PRIMARY = RGBColor(0x1B, 0x36, 0x5D)    # Azul marino institucional
    COLOR_SECONDARY = RGBColor(0x4A, 0x55, 0x68)  # Gris pizarra
    COLOR_MUTED = RGBColor(0x71, 0x80, 0x96)      # Gris claro
    HEX_HEADER_BG = "1B365D"
    HEX_ALT_BG = "F7FAFC"
    HEX_CALLOUT_BG = "F8FAFC"
    HEX_TOTAL_BG = "EDF2F7"

    # ================= 1. ENCABEZADO INSTITUCIONAL =================
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_before = Pt(0)
    p_header.paragraph_format.space_after = Pt(2)
    run_inst = p_header.add_run("UNIVERSIDAD DE GUADALAJARA\n")
    run_inst.font.name = 'Calibri'
    run_inst.font.size = Pt(13)
    run_inst.font.bold = True
    run_inst.font.color.rgb = COLOR_PRIMARY

    run_dept = p_header.add_run("Centro Universitario de Tonalá · Licenciatura en Ingeniería en Computación\nAdministración de Proyectos de Software · Cátedra: Dr. Gabriel Navarro Salcedo")
    run_dept.font.name = 'Calibri'
    run_dept.font.size = Pt(9.5)
    run_dept.font.color.rgb = COLOR_SECONDARY

    # Título principal
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(2)
    run_t1 = p_title.add_run("Actividad 2.3: Tabla comparativa de estimación del proyecto")
    run_t1.font.name = 'Calibri'
    run_t1.font.size = Pt(16)
    run_t1.font.bold = True
    run_t1.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(10)
    run_sub = p_sub.add_run("Evaluación y contraste de metodologías: Presupuesto inicial (WBS), COCOMO básico y COCOMO intermedio")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Tabla de Metadatos
    meta_table = doc.add_table(rows=2, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, color="CBD5E1", sz="4")
    
    meta_data = [
        [("Proyecto:", "Monchis Café (Web & POS)"), ("Alumno:", "Ernesto Fierro Meléndez"), ("Docente:", "Dr. Gabriel Navarro Salcedo"), ("Ciclo / Ponderación:", "2026-B · 2 Puntos (2%)")],
        [("Tamaño de software:", "22.5 KLOC (descomposición)"), ("Modo COCOMO:", "Semiacoplado (Semidetached)"), ("Horizonte planeado:", "16 semanas (4.0 meses)"), ("Fecha de entrega:", "Octubre de 2026")]
    ]

    for r_idx, row in enumerate(meta_data):
        for c_idx, (label, val) in enumerate(row):
            cell = meta_table.rows[r_idx].cells[c_idx]
            set_cell_shading(cell, HEX_CALLOUT_BG)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r_l = p.add_run(f"{label} ")
            r_l.font.size = Pt(8.5)
            r_l.font.bold = True
            r_l.font.color.rgb = COLOR_PRIMARY
            r_v = p.add_run(val)
            r_v.font.size = Pt(8.5)
            r_v.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Llamada de Objetivo
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(4)
    p_obj.paragraph_format.space_after = Pt(10)
    p_obj.paragraph_format.left_indent = Inches(0.2)
    p_obj.paragraph_format.right_indent = Inches(0.2)
    r_obj_title = p_obj.add_run("Objetivo general: ")
    r_obj_title.font.bold = True
    r_obj_title.font.size = Pt(9.5)
    r_obj_title.font.color.rgb = COLOR_PRIMARY
    r_obj_body = p_obj.add_run(
        "Evaluar y comparar cuantitativa y técnicamente tres métodos de estimación aplicados al proyecto Monchis Café: "
        "(1) el presupuesto inicial analítico (WBS / Bottom-Up) elaborado sobre la estructura de propuesta de presupuesto de Smartsheet, "
        "(2) el modelo paramétrico COCOMO básico (Boehm, 1981; Navarro Salcedo, 2026), y "
        "(3) el modelo COCOMO intermedio con 15 conductores de coste. "
        "El análisis concluye con la selección fundamentada del método más idóneo para la viabilidad contractual y de ingeniería del sistema."
    )
    r_obj_body.font.size = Pt(9.5)

    # Helper para encabezados de sección
    def add_section_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return p

    def add_section_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY
        return p

    # ================= SECCIÓN 1 =================
    add_section_h1("1. Contexto del proyecto y dimensionamiento del software")
    
    p_desc = doc.add_paragraph(
        "Monchis Café es una plataforma web integral de comercio electrónico y Punto de Venta (POS) en barra diseñada para una cafetería "
        "de especialidad enfocada en productos orgánicos de comercio justo. El sistema opera sobre una arquitectura monorepo distribuida "
        "y orientada a eventos para garantizar la resiliencia en los cobros y pedidos de mostrador:"
    )
    p_desc.paragraph_format.space_after = Pt(4)

    bullets_mod = [
        ("Frontend Web y POS de mostrador: ", "Desarrollado en Vue 3 con Pinia, prerenderizado estático mediante vite-ssg para optimización SEO en catálogo público, módulo de personalización de bebidas y panel táctil ágil de caja."),
        ("Backend y capa transaccional: ", "API modular en Next.js y Node.js con autenticación stateless mediante JWT rotativo en cookies httpOnly seguras, autorización por roles (RBAC) y persistencia relacional con PostgreSQL vía Prisma ORM."),
        ("Orquestación distribuida y mensajería: ", "Gestión asíncrona de órdenes y transacciones mediante el Patrón Saga orquestado en RabbitMQ (topic exchange cafeteria.events), incorporando colas con reintentos progresivos y colas de descarte (Dead Letter Queue - DLQ)."),
        ("Módulo de fidelización y trazabilidad: ", "Pasaporte digital QR para consultar el origen y lote de las fincas cafetaleras asociadas y control de beneficios por vaso reutilizable (-$5 MXN).")
    ]
    for b_title, b_desc in bullets_mod:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        r_bt = bp.add_run(b_title)
        r_bt.font.bold = True
        bp.add_run(b_desc)

    add_section_h2("1.1. Justificación del tamaño del software (22.5 KLOC)")
    p_kloc = doc.add_paragraph(
        "El tamaño estimado para el desarrollo nuevo del sistema se estableció en 22,500 líneas de código fuente (22.5 KLOC). "
        "Esta cifra se determinó mediante la técnica formal de descomposición por módulos (Pressman & Maxim, 2020), "
        "contrastando los componentes funcionales requeridos con la densidad de código promedio esperada en el stack seleccionado:"
    )
    p_kloc.paragraph_format.space_after = Pt(4)

    kloc_items = [
        ("Frontend (Vue 3, Pinia, vite-ssg) — ~8.5 KLOC: ", "Comprende unas 25 vistas y componentes interactivos (panel táctil de caja POS, selector interactivo de bebidas con siropes/leches, catálogo comercial, carrito de compra, resumen de caja y store global de estado en Pinia). Equivale a un promedio de ~340 líneas por componente/vista."),
        ("Backend y servicios de negocio (Next.js, TypeScript) — ~7.0 KLOC: ", "Comprende unos 18 endpoints REST y route handlers (procesamiento de pedidos, cobros, inventario, usuarios, autenticación y auditoría), middlewares de seguridad y esquemas de validación estricta con Zod. Equivale a un promedio de ~390 líneas por endpoint/servicio."),
        ("Persistencia y mensajería (Prisma ORM, RabbitMQ) — ~3.5 KLOC: ", "Modelado relacional en PostgreSQL con Prisma Schema, scripts de migración, seeders y la implementación del Patrón Saga (productores de eventos, consumidores, colas de retardo y reversas automáticas ante falta de insumos)."),
        ("Pruebas automatizadas y DevOps — ~3.5 KLOC: ", "Baterías de pruebas unitarias con Vitest, pruebas de integración de API, flujos end-to-end críticos con Playwright, Dockerfiles multicapa y configuración de workflows de CI/CD.")
    ]
    for kt, kd in kloc_items:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        r_kt = bp.add_run(kt)
        r_kt.font.bold = True
        bp.add_run(kd)

    p_kloc_note = doc.add_paragraph(
        "Nota sobre el alcance de KLOC: En la literatura clásica de COCOMO se consideran estrictamente líneas de código entregable. "
        "Si el equipo o el auditor optaran por excluir las 3.5 KLOC de pruebas automatizadas y scripts de infraestructura DevOps, "
        "el tamaño base del aplicativo central se situaría en 19.0 KLOC entregables."
    )
    p_kloc_note.paragraph_format.space_before = Pt(2)
    p_kloc_note.paragraph_format.space_after = Pt(6)
    p_kloc_note.runs[0].font.size = Pt(9)
    p_kloc_note.runs[0].font.italic = True
    p_kloc_note.runs[0].font.color.rgb = COLOR_MUTED

    # ================= SECCIÓN 2 =================
    add_section_h1("2. Método 1: Presupuesto inicial (WBS / Bottom-Up)")
    p_wbs_desc = doc.add_paragraph(
        "El presupuesto formal del proyecto se diseñó utilizando como referencia la estructura de propuesta de presupuesto de proyectos "
        "de Smartsheet (Smartsheet Inc., s. f.), siguiendo los estándares de descomposición analítica del trabajo del "
        "Project Management Institute (PMI, 2019). Este método ascendente (Bottom-Up) parte del detalle de actividades y recursos para "
        "construir el costo integral:"
    )
    p_wbs_desc.paragraph_format.space_after = Pt(4)

    wbs_table = doc.add_table(rows=6, cols=4)
    wbs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(wbs_table, color="1B365D", sz="4")

    headers_wbs = ["WBS", "Categoría / Paquete de trabajo", "Descripción operativa y alcance", "Monto total (MXN)"]
    for idx, h in enumerate(headers_wbs):
        cell = wbs_table.rows[0].cells[idx]
        set_cell_shading(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if idx == 3:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    wbs_rows = [
        ("1.0", "Mano de obra (Labor / Equipo técnico)", "5 roles técnicos (Scrum Master, Arquitecto/DevOps, Desarrollador Backend, Desarrollador Frontend, QA/Seguridad) cubriendo 9 paquetes de trabajo y 2,720 horas de desarrollo planeadas en 16 semanas.", "$358,000.00 MXN"),
        ("2.0", "Equipamiento y hardware de barra", "Terminal táctil Android 10.5\" para mostrador ($6,500), Impresora térmica de tickets de 80mm ($1,800) y Lector de códigos 1D/2D QR ($1,200).", "$9,500.00 MXN"),
        ("3.0", "Infraestructura Cloud y servicios SaaS", "Servicios administrados en Google Cloud Platform (Cloud Run, Cloud SQL, Redis: $7,200), CloudAMQP RabbitMQ ($2,200), Dominio y certificados TLS ($950) y licencias GitHub ($2,400) durante 4 meses.", "$12,750.00 MXN"),
        ("4.0", "Fondo de contingencia y mitigación (10%)", "Reserva del 10% sobre costos directos para imprevistos técnicos, fluctuaciones cambiarias o incrementos de consumo en la nube.", "$38,025.00 MXN"),
        ("TOTAL", "PRESUPUESTO TOTAL INICIAL", "Inversión integral estimada para el desarrollo, despliegue operativo y estabilización del sistema.", "$418,275.00 MXN")
    ]

    for r_idx, (wbs_code, cat, desc, amount) in enumerate(wbs_rows, start=1):
        row = wbs_table.rows[r_idx]
        is_total = (wbs_code == "TOTAL")
        bg = HEX_TOTAL_BG if is_total else (HEX_ALT_BG if r_idx % 2 == 1 else "FFFFFF")
        
        for c_idx, val in enumerate([wbs_code, cat, desc, amount]):
            cell = row.cells[c_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if is_total:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            if c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Métricas clave del presupuesto
    p_wbs_metrics = doc.add_paragraph()
    p_wbs_metrics.paragraph_format.space_before = Pt(6)
    p_wbs_metrics.paragraph_format.space_after = Pt(6)
    p_wbs_metrics.paragraph_format.left_indent = Inches(0.15)
    
    metrics_text = (
        "Métricas de ingeniería del presupuesto inicial:\n"
        "• Duración planificada: 16 semanas (4.0 meses de calendario en 8 sprints quincenales).\n"
        "• Esfuerzo equivalente en personas-mes: PM = 2,720 horas / 160 hrs/mes = 17.00 PM.\n"
        "• Personal promedio equivalente: NP = 17.0 PM / 4.0 meses = 4.25 personas a tiempo completo.\n"
        "• Costo promedio por hora / persona-mes: Tarifa promedio = $131.62 MXN/hr → Costo por PM = 160 hrs × $131.62 = $21,058.82 MXN/PM.\n"
        "• Productividad implícita: 22,500 LOC / 17.0 PM = 1,323.53 LOC / Persona-Mes (~66 LOC diarias por persona).\n"
        "• Análisis de riesgo de productividad: Esta tasa de 1,323 LOC/PM se fundamenta en la reutilización de librerías, tipado estricto y componentes preconstruidos. "
        "Sin embargo, constituye una meta agresiva. Si surgen bloqueos en la integración con el hardware de mostrador o si el alcance supera las 22.5 KLOC, "
        "el cronograma de 16 semanas enfrentará riesgos serios de desviación."
    )
    r_wm = p_wbs_metrics.add_run(metrics_text)
    r_wm.font.size = Pt(9)
    r_wm.font.color.rgb = COLOR_SECONDARY

    # ================= SECCIÓN 3 =================
    add_section_h1("3. Método 2: Estimación por COCOMO básico (Boehm, 1981)")
    p_cocomo_b_desc = doc.add_paragraph(
        "El modelo COCOMO básico (Constructive Cost Model) calcula el esfuerzo y el calendario como funciones de potencia empíricas "
        "dependientes exclusivamente del tamaño del software en miles de líneas de código fuente (KLOC; Boehm, 1981). "
        "Para Monchis Café (22.5 KLOC, arquitectura distribuida y equipo con experiencia intermedia), la clasificación corresponde "
        "al Modo Semiacoplado (Semidetached) (Navarro Salcedo, 2026):"
    )
    p_cocomo_b_desc.paragraph_format.space_after = Pt(4)

    # Ecuaciones
    p_eq = doc.add_paragraph()
    p_eq.paragraph_format.left_indent = Inches(0.2)
    p_eq.paragraph_format.space_after = Pt(6)
    eq_text = (
        "Ecuaciones del modelo semiacoplado (A = 3.00, B = 1.12, C = 2.50, D = 0.35):\n"
        "1. Esfuerzo nominal: PM = A · (KLOC)^B = 3.00 · (22.5)^1.12\n"
        "2. Tiempo de desarrollo: TDEV = C · (PM)^D = 2.50 · (PM)^0.35\n"
        "3. Personal promedio requerido: NP = PM / TDEV\n"
        "4. Costo de mano de obra estimado: Costo = PM · Tarifa promedio mensual"
    )
    r_eq = p_eq.add_run(eq_text)
    r_eq.font.size = Pt(9)
    r_eq.font.bold = True
    r_eq.font.color.rgb = COLOR_PRIMARY

    add_section_h2("3.1. Desarrollo matemático paso a paso")
    p_calc_b = doc.add_paragraph()
    p_calc_b.paragraph_format.left_indent = Inches(0.2)
    p_calc_b.paragraph_format.space_after = Pt(6)
    calc_b_text = (
        "• Cálculo del esfuerzo (PM):\n"
        "  (22.5)^1.12 = exp(1.12 · ln(22.5)) = exp(1.12 · 3.113515) = exp(3.487137) = 32.6922\n"
        "  PM = 3.00 × 32.6922 = 98.0767 PM ≈ 98.08 Personas-Mes.\n\n"
        "• Cálculo del tiempo de desarrollo (TDEV):\n"
        "  (98.0767)^0.35 = exp(0.35 · ln(98.0767)) = exp(0.35 · 4.5857) = exp(1.6050) = 4.9779\n"
        "  TDEV = 2.50 × 4.9779 = 12.4448 meses ≈ 12.44 Meses de calendario (~54 semanas).\n\n"
        "• Personal promedio (NP) y productividad:\n"
        "  NP = 98.0767 / 12.4448 = 7.88 personas a tiempo completo.\n"
        "  Productividad nominal = 22,500 LOC / 98.08 PM = 229.41 LOC / Persona-Mes.\n\n"
        "• Costo de mano de obra estimado (Tarifa base: $21,058.82 MXN/PM):\n"
        "  Costo = 98.0767 PM × $21,058.82 MXN/PM = $2,065,379.11 MXN."
    )
    r_cb = p_calc_b.add_run(calc_b_text)
    r_cb.font.size = Pt(9)
    r_cb.font.color.rgb = COLOR_SECONDARY

    p_crit_b = doc.add_paragraph(
        "Evaluación del modelo básico: COCOMO básico arroja 98.08 personas-mes y un plazo de 12.44 meses. "
        "Este modelo fue concebido bajo el paradigma procedural de 1981, donde no existían frameworks de alto nivel ni reutilización masiva de librerías, "
        "lo que conduce a una sobreestimación notable del esfuerzo para una aplicación web moderna. Asimismo, pretender incorporar casi 8 desarrolladores "
        "a tiempo completo en un proyecto de 22.5 KLOC provocaría una sobrecarga en la comunicación interna, en consonancia con la ley de Brooks (1975)."
    )
    p_crit_b.paragraph_format.space_after = Pt(8)

    # ================= SECCIÓN 4 =================
    add_section_h1("4. Método 3: Estimación por COCOMO intermedio")
    p_cocomo_i_desc = doc.add_paragraph(
        "COCOMO intermedio refina la estimación básica incorporando el Factor de Ajuste de Esfuerzo (EAF - Effort Adjustment Factor), "
        "obtenido como el producto de 15 conductores de coste (Cost Drivers; Boehm, 1981) evaluados de acuerdo con el contexto técnico del proyecto:"
    )
    p_cocomo_i_desc.paragraph_format.space_after = Pt(4)

    drivers_table = doc.add_table(rows=17, cols=5)
    drivers_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(drivers_table, color="1B365D", sz="4")

    headers_drivers = ["Categoría", "Conductor de coste (Cost Driver)", "Nivel", "Valor", "Justificación técnica en Monchis Café"]
    for idx, h in enumerate(headers_drivers):
        cell = drivers_table.rows[0].cells[idx]
        set_cell_shading(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if idx in [2, 3]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    drivers_data = [
        ("Producto", "RELY — Fiabilidad requerida del software", "Alta", "1.15", "Manejo de pedidos, cobros POS en barra y persistencia transaccional de ventas."),
        ("Producto", "DATA — Tamaño de la base de datos", "Nominal", "1.00", "Base de datos relacional PostgreSQL con volumen de datos comercial moderado."),
        ("Producto", "CPLX — Complejidad del software", "Alta", "1.15", "Patrón Saga distribuido con RabbitMQ, transacciones y autenticación stateless."),
        ("Plataforma", "TIME — Restricción de tiempo de ejecución", "Nominal", "1.00", "Aplicación transaccional web sin restricciones de tiempo real estricto."),
        ("Plataforma", "STOR — Restricción de memoria principal", "Nominal", "1.00", "Memoria holgada en instancias serverless de Cloud Run y contenedores Docker."),
        ("Plataforma", "VIRT — Volatilidad de la máquina virtual", "Baja", "0.87", "Entorno estandarizado en contenedores Docker y Node.js LTS de alta estabilidad."),
        ("Plataforma", "TURN — Tiempo de respuesta de la máquina", "Nominal", "1.00", "Entorno de desarrollo ágil con recarga en caliente (HMR) mediante Vite."),
        ("Personal", "ACAP — Capacidad de los analistas", "Alta", "0.86", "Equipo con dominio de los requerimientos y arquitectura del sistema."),
        ("Personal", "AEXP — Experiencia en la aplicación", "Nominal", "1.00", "Experiencia media en aplicaciones de comercio electrónico y puntos de venta."),
        ("Personal", "PCAP — Capacidad de los programadores", "Alta", "0.86", "Dominio sólido en TypeScript, Vue 3, Next.js y modelado Prisma."),
        ("Personal", "VEXP — Experiencia en máquina virtual", "Nominal", "1.00", "Familiaridad estándar con entornos Linux y contenedores Docker."),
        ("Personal", "LEXP — Experiencia en lenguajes de prog.", "Alta", "0.95", "Especialización en TypeScript y desarrollo full-stack en JavaScript moderno."),
        ("Proyecto", "MODP — Prácticas modernas de programación", "Muy Alta", "0.82", "Uso intensivo de arquitectura monorepo Turborepo, TDD, linting y CI/CD."),
        ("Proyecto", "TOOL — Uso de herramientas de software", "Alta", "0.91", "Herramientas avanzadas: Prisma Studio, Playwright, Vitest y RabbitMQ UI."),
        ("Proyecto", "SCED — Plazo de desarrollo requerido", "Nominal*", "1.00", "Se evalúa nominal (1.00) para aislar tecnología y equipo (ver análisis técnico)."),
        ("TOTAL", "PRODUCTO DE LOS 15 FACTORES (EAF)", "-", "0.6032", "Factor de Ajuste de Esfuerzo neto (representa un ahorro del 39.7% frente al nominal).")
    ]

    for r_idx, (cat, d_name, lvl, val, just) in enumerate(drivers_data, start=1):
        row = drivers_table.rows[r_idx]
        is_total = (cat == "TOTAL")
        bg = HEX_TOTAL_BG if is_total else (HEX_ALT_BG if r_idx % 2 == 1 else "FFFFFF")
        
        for c_idx, text_val in enumerate([cat, d_name, lvl, val, just]):
            cell = row.cells[c_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text_val)
            r.font.size = Pt(8)
            if is_total:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            if c_idx in [2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_section_h2("4.1. Cálculos de COCOMO intermedio y análisis de SCED")
    p_calc_i = doc.add_paragraph()
    p_calc_i.paragraph_format.left_indent = Inches(0.2)
    p_calc_i.paragraph_format.space_after = Pt(6)
    calc_i_text = (
        "• Cálculo aritmético del factor EAF:\n"
        "  EAF = 1.15 × 1.00 × 1.15 × 1.00 × 1.00 × 0.87 × 1.00 × 0.86 × 1.00 × 0.86 × 1.00 × 0.95 × 0.82 × 0.91 × 1.00 = 0.6032\n\n"
        "• Esfuerzo ajustado (PM_ajustado):\n"
        "  PM_ajustado = EAF × PM_nominal = 0.60324 × 98.0767 = 59.1638 PM ≈ 59.16 Personas-Mes.\n"
        "  (Reducción de 38.91 personas-mes respecto al modelo básico, ahorro del 39.7%).\n\n"
        "• Tiempo de desarrollo ajustado (TDEV_ajustado):\n"
        "  (59.1638)^0.35 = 4.1708\n"
        "  TDEV_ajustado = 2.50 × 4.1708 = 10.4270 meses ≈ 10.43 Meses de calendario (~45 semanas).\n\n"
        "• Personal promedio (NP) y productividad ajustada:\n"
        "  NP = 59.1638 / 10.4270 = 5.67 personas a tiempo completo.\n"
        "  Productividad ajustada = 22,500 LOC / 59.16 PM = 380.30 LOC / Persona-Mes (+65.8% vs. básico).\n\n"
        "• Costo de mano de obra ajustado (Tarifa base: $21,058.82 MXN/PM):\n"
        "  Costo = 59.1638 PM × $21,058.82 MXN/PM = $1,245,920.89 MXN."
    )
    r_ci = p_calc_i.add_run(calc_i_text)
    r_ci.font.size = Pt(9)
    r_ci.font.color.rgb = COLOR_SECONDARY

    p_sced_note = doc.add_paragraph(
        "*Análisis técnico sobre la restricción de plazo (SCED): En COCOMO 81, el tiempo nominal estimado para este proyecto es de 10.43 meses. "
        "Desarrollar el sistema en 4 meses (16 semanas) representaría comprimir el calendario al ~38% del plazo nominal. "
        "Barry Boehm (1981) estableció explícitamente que comprimir un cronograma por debajo del 75% del plazo nominal no es alcanzable dentro de las dinámicas "
        "de desarrollo modeladas por COCOMO. Si en el ejercicio se forzara el factor de compresión máxima disponible en las tablas (SCED = Muy Bajo, multiplicador 1.23), "
        "el EAF subiría a 0.7420 y el esfuerzo estimado a 72.77 PM, con un costo aproximado de ≈ $1.53 M MXN. "
        "Aun así, dado que 4 meses es una compresión mucho más agresiva que el 75%, el multiplicador 1.23 en realidad subestima la penalización que el modelo clásico "
        "asignaría. Por ello, se mantiene SCED nominal (1.00) en el cálculo base con el fin de aislar y evaluar objetivamente el efecto exclusivo de las herramientas "
        "y el perfil técnico del equipo sobre el esfuerzo."
    )
    p_sced_note.paragraph_format.space_after = Pt(8)
    p_sced_note.runs[0].font.size = Pt(8.8)
    p_sced_note.runs[0].font.italic = True
    p_sced_note.runs[0].font.color.rgb = COLOR_SECONDARY

    # ================= SECCIÓN 5 =================
    doc.add_page_break()
    add_section_h1("5. Tabla comparativa de las tres estimaciones")
    p_comp_intro = doc.add_paragraph(
        "A continuación se presenta la matriz comparativa que contrasta las dimensiones operativas, económicas y de ingeniería "
        "entre el presupuesto inicial ascendente y los modelos paramétricos de Boehm:"
    )
    p_comp_intro.paragraph_format.space_after = Pt(4)

    comp_table = doc.add_table(rows=16, cols=4)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(comp_table, color="1B365D", sz="4")

    headers_comp = ["Dimensión / Criterio", "1. Presupuesto inicial (WBS)", "2. COCOMO básico (Boehm, 1981)", "3. COCOMO intermedio (Boehm, 1981)"]
    for idx, h in enumerate(headers_comp):
        cell = comp_table.rows[0].cells[idx]
        set_cell_shading(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if idx > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    comp_data = [
        ("Filosofía y enfoque", "Bottom-Up (Ascendente): Descomposición analítica en paquetes de trabajo y asignación de horas/perfiles (PMI, 2019).", "Paramétrico: Modelo empírico basado estrictamente en volumen de líneas de código fuente (KLOC; Boehm, 1981).", "Paramétrico calibrado: Modelo empírico modificado por 15 conductores de coste de contexto (EAF; Boehm, 1981)."),
        ("Variable primaria de entrada", "Horas por rol técnico, requerimientos de hardware, servicios Cloud y reserva de riesgos.", "Tamaño del software en KLOC (22.5 KLOC) y modo de desarrollo (Semiacoplado).", "Tamaño (22.5 KLOC), modo Semiacoplado y vector de 15 factores de coste (EAF = 0.6032)."),
        ("Esfuerzo estimado (PM)", "17.00 Personas-Mes (2,720 horas de desarrollo planeadas).", "98.08 Personas-Mes (+476.9% respecto al plan WBS).", "59.16 Personas-Mes (Ahorro del 39.7% frente al básico)."),
        ("Plazo de calendario (TDEV)", "4.00 Meses (16 semanas / 8 sprints quincenales).", "12.44 Meses (~54 semanas de calendario).", "10.43 Meses (~45 semanas de calendario)."),
        ("Personal promedio (NP)", "4.25 Desarrolladores (Equipo planeado: 5 roles técnicos).", "7.88 Desarrolladores (Riesgo de sobrecarga comunicativa; Brooks, 1975).", "5.67 Desarrolladores (Equipo mediano estructurado)."),
        ("Productividad proyectada", "1,323.53 LOC / PM (Alta productividad; meta exigente sujeta a riesgos).", "229.41 LOC / PM (Típica de desarrollo procedural en C/Fortran).", "380.30 LOC / PM (+65.8% de productividad respecto al modelo básico)."),
        ("Costo mano de obra (Labor)", "$358,000.00 MXN (Tarifas de mercado en la ZMG: $112.50 a $175.00 MXN/hr).", "$2,065,379.11 MXN (Valorizado a tarifa calculada de $21,058.82 MXN/PM).", "$1,245,920.89 MXN (Valorizado a tarifa calculada de $21,058.82 MXN/PM)."),
        ("Equipamiento de mostrador / POS", "$9,500.00 MXN (Tablet táctil, impresora térmica de tickets y escáner QR).", "$0.00 (Omitido; no contempla activos físicos de hardware).", "$0.00 (Omitido; no contempla activos físicos de hardware)."),
        ("Infraestructura Cloud y licencias", "$12,750.00 MXN (GCP Cloud Run, SQL, RabbitMQ CloudAMQP, SSL y GitHub).", "$0.00 (Omitido; no contempla servidores ni servicios en la nube).", "$0.00 (Omitido; no contempla servidores ni servicios en la nube)."),
        ("Fondo de contingencia", "$38,025.00 MXN (10% explícito sobre costos directos para mitigación).", "$0.00 (Omitido; no incluye reserva financiera de imprevistos).", "$0.00 (Omitido; no incluye reserva financiera de imprevistos)."),
        ("INVERSIÓN TOTAL ESTIMADA", "$418,275.00 MXN\n(Mano de obra, hardware, nube y contingencia)", "$2,065,379.11 MXN\n(Solo mano de obra estimada)", "$1,245,920.89 MXN\n(Solo mano de obra estimada)"),
        ("Sensibilidad a tecnología moderna", "Alta: El desglose de horas refleja directamente la velocidad de TypeScript, Prisma y Vue.", "Nula: Modela una línea de código moderno igual que en 1981.", "Moderada: Modula el impacto mediante MODP (0.82), TOOL (0.91), PCAP (0.86) y ACAP (0.86)."),
        ("Sensibilidad a complejidad", "Específica por tarea (ej. 480 hrs en pruebas/seguridad; 400 hrs en POS).", "Solo reflejada en el exponente B = 1.12 del modo semiacoplado.", "Ponderada explícitamente en RELY (1.15) y CPLX (1.15), compensada por MODP."),
        ("Principales ventajas", "• Control presupuestal integral.\n• Incluye hardware, nube y riesgos.\n• Aceptado contractualmente por clientes.", "• Rápido y sencillo de calcular.\n• Solo requiere estimar KLOC.\n• Referencia histórica macroscópica.", "• Calibra el perfil del equipo.\n• Premia buenas prácticas modernas.\n• Reduce drásticamente la sobreestimación."),
        ("Principales limitaciones", "• Vulnerable al sesgo de optimismo.\n• Requiere descomposición detallada previa.\n• No provee fórmulas de compresión.", "• Sobrestima severamente en entornos web.\n• Ignora herramientas y frameworks.\n• Desconectado de presupuestos comerciales.", "• Sigue omitiendo costos de nube y equipo.\n• Calibrado con proyectos de 1981.\n• Requiere justificar 15 coeficientes.")
    ]

    for r_idx, (crit, c_wbs, c_bas, c_int) in enumerate(comp_data, start=1):
        row = comp_table.rows[r_idx]
        is_total = (crit == "INVERSIÓN TOTAL ESTIMADA")
        bg = HEX_TOTAL_BG if is_total else (HEX_ALT_BG if r_idx % 2 == 1 else "FFFFFF")
        
        for c_idx, val in enumerate([crit, c_wbs, c_bas, c_int]):
            cell = row.cells[c_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(7.8)
            if is_total:
                r.font.bold = True
                r.font.color.rgb = COLOR_PRIMARY

    # ================= SECCIÓN 6 =================
    add_section_h1("6. Análisis técnico comparativo y conclusiones")
    
    add_section_h2("6.1. Justificación de las discrepancias entre modelos (1981 vs. 2026)")
    p_disc = doc.add_paragraph(
        "Al contrastar el presupuesto inicial ($418,275.00 MXN | 17.00 PM) con los modelos COCOMO básico ($2,065,379.11 MXN | 98.08 PM) "
        "e intermedio ($1,245,920.89 MXN | 59.16 PM), se hace evidente una brecha técnica sustancial: el esfuerzo del modelo básico "
        "representa casi seis veces el esfuerzo planeado en el presupuesto (+476.9%), mientras que el modelo intermedio representa "
        "más de tres veces dicho esfuerzo (59.16 vs. 17.00 PM, un +248%). Esta divergencia se explica por los siguientes factores:"
    )
    p_disc.paragraph_format.space_after = Pt(4)

    reasons = [
        ("Base empírica de COCOMO 81: ", "Barry Boehm construyó el modelo original en 1981 analizando 63 proyectos tradicionales de software en FORTRAN, COBOL y ensamblador (Boehm, 1981). En aquel contexto, escribir 22.5 KLOC demandaba construir manualmente drivers de comunicación, interfaces de texto y rutinas de persistencia a bajo nivel, lo que explica la proyección de 98 personas-mes."),
        ("Ecosistema de desarrollo moderno: ", "En Monchis Café, una línea de código en TypeScript moderno posee una densidad funcional significativamente más alta: se emplean componentes reactivos en Vue 3, esquemas automáticos con Prisma ORM y servicios administrados en la nube. Esto permite proyectar una productividad de 1,323 LOC/PM en 16 semanas. Sin embargo, debe reconocerse que este rendimiento asume un flujo de trabajo continuo y sin fricciones mayores; si surgen imprevistos en la integración de periféricos o el código crece, el cronograma corre el riesgo de desfasarse."),
        ("Efecto de los conductores de coste: ", "Al incorporar los 15 conductores en COCOMO intermedio (EAF = 0.6032), el esfuerzo disminuye de 98 a 59 personas-mes. Los factores determinantes en esta reducción son las prácticas modernas (MODP = 0.82) y la capacidad del equipo (ACAP y PCAP = 0.86). Con todo, el cálculo resultante continúa situándose muy por encima del presupuesto inicial, y los 10.4 meses estimados distan de las 16 semanas contempladas.")
    ]
    for r_title, r_desc in reasons:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(2)
        r_rt = bp.add_run(r_title)
        r_rt.font.bold = True
        bp.add_run(r_desc)

    add_section_h2("6.2. Conclusión: Selección y recomendación metodológica")
    p_concl = doc.add_paragraph(
        "A partir de la evaluación técnica realizada, se establecen las siguientes conclusiones para la gestión de Monchis Café:\n"
        "1. El presupuesto inicial (WBS / Bottom-Up) es el método más adecuado para la viabilidad financiera y contractual: "
        "Es la única alternativa que ofrece una visión integral del negocio, ya que incorpora los $9,500 MXN de equipamiento de barra indispensable para operar en mostrador, "
        "los $12,750 MXN de servicios Cloud para alojar la base de datos y RabbitMQ, y el fondo de contingencia del 10% ($38,025 MXN). "
        "Los modelos COCOMO no contemplan estos costos directos, pues están estructurados exclusivamente para estimar esfuerzo de desarrollo de software.\n"
        "2. COCOMO intermedio: mejor referencia que el Básico, pero requiere calibración: "
        "COCOMO intermedio supera con claridad al modelo básico porque modula el esfuerzo en función del perfil técnico y las herramientas. "
        "No obstante, no debe emplearse como presupuesto comercial vinculante para el cliente, sino como una herramienta de contraste y auditoría cruzada "
        "que previene al equipo sobre el riesgo de sobrecarga comunicativa interna (Brooks, 1975) si se pretendiera resolver retrasos incorporando personal a destiempo. "
        "Para proyectos contemporáneos, la disciplina recurre a evoluciones como COCOMO II (Boehm et al., 2000), diseñado para desarrollo moderno basado en componentes "
        "y reutilización de software."
    )
    p_concl.paragraph_format.space_after = Pt(6)

    # 6.3 Reflexión Personal (dejada vacía con pauta para que el estudiante la redacte)
    add_section_h2("6.3. Reflexión personal sobre el proceso de estimación")
    p_ref_empty = doc.add_paragraph()
    p_ref_empty.paragraph_format.space_after = Pt(12)
    r_empty_note = p_ref_empty.add_run(
        "[Espacio reservado para la reflexión personal del alumno sobre el proceso de estimación: "
        "redacta aquí tus impresiones directas sobre las decisiones tomadas, los paquetes de trabajo que mayor complejidad implicaron "
        "—como la integración de periféricos POS o la orquestación con RabbitMQ— y tu aprendizaje al contrastar las fórmulas teóricas de 1981 con el presupuesto WBS.]\n\n\n\n"
    )
    r_empty_note.font.italic = True
    r_empty_note.font.color.rgb = COLOR_MUTED

    # ================= SECCIÓN 7: REFERENCIAS =================
    add_section_h1("7. Referencias bibliográficas (Normas APA 7)")
    refs = [
        "Boehm, B. W. (1981). Software engineering economics. Prentice-Hall.",
        "Boehm, B., Abts, C., Brown, A. W., Chulani, S., Clark, B. K., Horowitz, E., Madachy, R., Reifer, D., & Steece, B. (2000). Software cost estimation with COCOMO II. Prentice Hall.",
        "Brooks, F. P. (1975). The mythical man-month: Essays on software engineering. Addison-Wesley.",
        "Computrabajo México. (2026). Portal de empleo y sondeo de salarios en tecnología [Base de datos en línea]. https://www.computrabajo.com.mx/",
        "Glassdoor. (2026). Sueldos para desarrolladores de software en Guadalajara, Jalisco [Base de datos en línea]. https://www.glassdoor.com.mx/",
        "Navarro Salcedo, G. (2026). Material de apoyo y notas de clase sobre los modelos COCOMO básico e intermedio [Material didáctico de cátedra]. Universidad de Guadalajara.",
        "Pressman, R. S., & Maxim, B. R. (2020). Software engineering: A practitioner's approach (9.ª ed.). McGraw-Hill Education.",
        "Project Management Institute. (2019). Practice standard for work breakdown structures (3.ª ed.). Project Management Institute.",
        "Smartsheet Inc. (s. f.). Plantilla de propuesta de presupuesto de proyectos [Plantilla de hoja de cálculo]. Smartsheet. https://es.smartsheet.com/content/project-budget-templates"
    ]
    for rf in refs:
        p_rf = doc.add_paragraph()
        p_rf.paragraph_format.left_indent = Inches(0.4)
        p_rf.paragraph_format.first_line_indent = Inches(-0.4)
        p_rf.paragraph_format.space_after = Pt(3)
        r_rf = p_rf.add_run(rf)
        r_rf.font.size = Pt(8.5)
        r_rf.font.color.rgb = COLOR_SECONDARY

    docx_path = r"d:\ernestofm\Proyecto\docs\entregables\Actividad_2_3_Tabla_Comparativa_Estimacion_MonchisCafe.docx"
    downloads_docx = r"C:\Users\erfierro\Downloads\Actividad_2_3_Tabla_Comparativa_Estimacion_MonchisCafe.docx"

    doc.save(docx_path)
    print(f"Documento DOCX generado exitosamente en: {docx_path}")
    shutil.copy2(docx_path, downloads_docx)
    print(f"Copia creada exitosamente en Downloads: {downloads_docx}")

if __name__ == "__main__":
    create_styled_document()
