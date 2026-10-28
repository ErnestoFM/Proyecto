# -*- coding: utf-8 -*-
"""
Constructor del Libro Excel Oficial:
Actividad 4: Gestión de Finanzas — Monchis Café
Cátedra: Proyecto Integrador de Desarrollo de Software
Docente: Dr. Gabriel Navarro Salcedo
Estudiante: Ernesto Hatuey Fierro Meléndez
"""

import os
import glob
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_financial_model():
    pattern = r'C:\Users\erfierro\Downloads\*PLANTILLA PARA EL COSTO TOTAL*.xlsx'
    files = [f for f in glob.glob(pattern) if not f.endswith('~$')]
    if not files:
        fallback = r'd:\ernestofm\Proyecto\docs\presupuesto\Presupuesto_Monchis_Cafe.xlsx'
        if os.path.exists(fallback):
            excel_template_path = fallback
        else:
            raise FileNotFoundError("No se encontró el archivo de plantilla base.")
    else:
        excel_template_path = files[0]
        
    print(f"Cargando plantilla base: {excel_template_path}")
    wb = openpyxl.load_workbook(excel_template_path)
    
    # -------------------------------------------------------------
    # 1. PARAMETROS
    # -------------------------------------------------------------
    ws_params = wb['PARAMETROS']
    ws_params['B2'].value = "Monchis Café — Plataforma Web y POS Seguro para Cafetería Orgánica"
    ws_params['C3'].value = "Año de ejercicio fiscal (tarifas de retención de ISR y UMA 2026)"
    ws_params['B4'].value = 4      # 4 meses de ejecución
    ws_params['B5'].value = 1      # 1 proyecto a entregar
    ws_params['B6'].value = 0.30   # Margen 30%
    ws_params['B7'].value = 0.16   # IVA 16%
    ws_params['B8'].value = 0.10   # Contingencia 10%
    ws_params['B9'].value = 0.30   # Carga patronal 30%
    ws_params['C9'].value = "Aproximación técnica: IMSS, aportación INFONAVIT 5% e ISN local 3% (Jalisco/Michoacán)"
    ws_params['B10'].value = 160   # Horas al mes por persona
    ws_params['B14'].value = 0.18  # 18% interés anual préstamo
    ws_params['B15'].value = 24    # 24 meses plazo préstamo
    ws_params['A16'].value = "Fracción de inversión inicial imputada al proyecto"
    ws_params['B16'].value = 1.0   # 100% de depreciación imputada a los 4 meses
    
    # -------------------------------------------------------------
    # 2. NOMINA
    # -------------------------------------------------------------
    ws_nom = wb['NOMINA']
    ws_nom['A2'].value = "Líder de Proyecto / Scrum Master & DevOps"
    ws_nom['B2'].value = 36000.00
    ws_nom['C2'].value = 1
    
    ws_nom['A3'].value = "Desarrolladora Frontend Core (Vue 3 / UI)"
    ws_nom['B3'].value = 20000.00
    ws_nom['C3'].value = 1
    
    ws_nom['A4'].value = "Desarrollador Backend Core (Node.js / APIs)"
    ws_nom['B4'].value = 20000.00
    ws_nom['C4'].value = 1
    
    ws_nom['A7'].value = "Especialista en QA, Automatización & Ciberseguridad"
    ws_nom['B7'].value = 13500.00
    ws_nom['C7'].value = 1
    
    for r in [5, 6, 8, 9, 10, 11, 12, 13]:
        ws_nom[f'B{r}'].value = None
        ws_nom[f'C{r}'].value = None

    # -------------------------------------------------------------
    # 3. GASTOS FIJOS (7 rubros desglosados)
    # -------------------------------------------------------------
    ws_gf = wb['GASTOS FIJOS']
    gf_data = [
        (2, 6000.00, "Espacio de trabajo colaborativo / coworking y sala de reuniones de sprint"),
        (3, 1200.00, "Consumo eléctrico de estaciones de cómputo y climatización del espacio"),
        (4, 1500.00, "Enlace empresarial simétrico de fibra óptica 500 Mbps con IP fija para pruebas"),
        (5, 2500.00, "Servicios contables externos, dispersión de nómina y timbrado fiscal ante el SAT"),
        (6, 1800.00, "Herramientas SaaS de colaboración (Google Workspace, Slack Pro, Notion, Figma)"),
        (7, 1000.00, "Gestión comercial, demostraciones con el cliente y prospección de cuenta"),
        (8, 1000.00, "Mantenimiento preventivo, artículos de limpieza, papelería y consumibles de oficina"),
    ]
    for r, monto, crit in gf_data:
        ws_gf[f'B{r}'].value = monto
        ws_gf[f'E{r}'].value = crit
        
    # -------------------------------------------------------------
    # 4. COSTOS DIRECTOS (10 rubros reales = $26,800.00)
    # -------------------------------------------------------------
    ws_cd = wb['COSTOS DIRECTOS']
    cd_data = [
        (2, "Meses", 4, 1800.00, "Google Cloud Platform: Cloud Run (microservicios), Cloud SQL PostgreSQL, Redis Memorystore"),
        (3, "Meses", 4, 600.00, "Licencias de repositorios y minutos CI/CD en GitHub Team, y monitoreo Sentry"),
        (4, "Meses", 4, 550.00, "Broker de mensajería CloudAMQP RabbitMQ dedicado (Patrón Saga y Dead Letter Queue)"),
        (5, "Licencia anual", 1, 950.00, "Registro de dominio oficial corporativo .com.mx y certificado TLS/SSL Wildcard"),
        (6, "Dispositivos/Lic", 2, 1500.00, "Dispositivos móviles de laboratorio para testing multi-resolución, soporte táctil y escaneo QR"),
        (7, "Servicio", 1, 4500.00, "Subcontratación de pentesting ético externo y auditoría de pasarela de pago"),
        (8, "Visitas en sitio", 6, 500.00, "Viáticos para visitas de campo a cafetería en La Piedad (instalación física y pruebas en barra)"),
        (9, "Lote", 1, 850.00, "Materiales y consumibles (papel térmico 80mm sin BPA, cables de red y USB blindados)"),
        (10, "Servicio", 1, 1200.00, "Auditoría de código estático (SAST) y análisis de dependencias OWASP Top 10"),
        (11, "Paquete", 1, 1500.00, "Otros: Manuales impresos de mostrador y señalética de lealtad Monchis Rewards")
    ]
    for r, u, cant, cu, notas in cd_data:
        ws_cd[f'B{r}'].value = u
        ws_cd[f'C{r}'].value = cant
        ws_cd[f'D{r}'].value = cu
        ws_cd[f'F{r}'].value = notas

    # -------------------------------------------------------------
    # 5. INVERSION INICIAL (Depreciación proporcional en 4 meses)
    # -------------------------------------------------------------
    ws_inv = wb['INVERSION INICIAL']
    inv_data = [
        (2, 4, 22000.00, 36),  # Laptops (36 meses)
        (3, 4, 4500.00, 60),   # Mobiliario ergonómico (60 meses)
        (4, 1, 8500.00, 36),   # Equipo de red de laboratorio (36 meses)
        (5, 4, 3200.00, 24),   # Licencias IDEs y herramientas de desarrollo (24 meses)
        (6, 1, 12000.00, 12),  # Capacitación técnica en RabbitMQ y Playwright (12 meses)
        (7, 1, 9500.00, 24),   # Kit de mostrador POS de pruebas (24 meses)
    ]
    for r, cant, pu, vu in inv_data:
        ws_inv[f'B{r}'].value = cant
        ws_inv[f'C{r}'].value = pu
        ws_inv[f'E{r}'].value = vu

    # -------------------------------------------------------------
    # 6. PRESTAMO (Interés promedio prorrateado a 4 meses)
    # -------------------------------------------------------------
    ws_prest = wb['PRESTAMO']
    ws_prest['B2'].value = 40000.00

    # -------------------------------------------------------------
    # 7. HOJA REESTRUCTURADA: ESTRATEGIA DE VENTA
    # -------------------------------------------------------------
    sheet_name = 'ESTRATEGIA DE VENTA'
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    
    ws_strat = wb.create_sheet(title=sheet_name)
    ws_strat.views.sheetView[0].showGridLines = True
    
    NAVY_DARK = "17365D"
    NAVY_MED = "244B7A"
    BLUE_LIGHT = "D9E2EC"
    BLUE_PALE = "F0F4F8"
    GREEN_SUCCESS = "E2F0D9"
    GRAY_LIGHT = "F8FAFC"
    WHITE = "FFFFFF"
    
    font_main = "Arial"
    
    f_title = Font(name=font_main, size=14, bold=True, color=WHITE)
    f_subtitle = Font(name=font_main, size=10, italic=True, color=WHITE)
    f_section = Font(name=font_main, size=11, bold=True, color=WHITE)
    f_subhead = Font(name=font_main, size=10, bold=True, color=NAVY_DARK)
    f_tbl_hdr = Font(name=font_main, size=9.5, bold=True, color=WHITE)
    f_bold = Font(name=font_main, size=9.5, bold=True, color="000000")
    f_regular = Font(name=font_main, size=9.5, color="1E293B")
    f_italic = Font(name=font_main, size=8.5, italic=True, color="64748B")
    f_ref = Font(name="Consolas", size=8.5, color="2B4C7E")
    
    fill_header = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
    fill_section = PatternFill(start_color=NAVY_MED, end_color=NAVY_MED, fill_type="solid")
    fill_subhead = PatternFill(start_color=BLUE_LIGHT, end_color=BLUE_LIGHT, fill_type="solid")
    fill_zebra = PatternFill(start_color=BLUE_PALE, end_color=BLUE_PALE, fill_type="solid")
    fill_success = PatternFill(start_color=GREEN_SUCCESS, end_color=GREEN_SUCCESS, fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    box_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    
    fmt_currency = "$#,##0.00"
    fmt_pct = "0.0%"
    fmt_num = "#,##0"
    
    col_widths = {
        'A': 4,
        'B': 30,  # Concepto
        'C': 16,  # Valor
        'D': 24,  # Referencia Celda
        'E': 16,  # % s/ Venta o CTP
        'F': 22,  # Interpretación Financiera
        'G': 46,  # Criterio y Justificación
        'H': 4
    }
    for col, width in col_widths.items():
        ws_strat.column_dimensions[col].width = width
        
    curr_row = 2
    
    def apply_merged_borders(start_col, start_r, end_col, end_r, fill=None, font=None, alignment=None):
        for r_idx in range(start_r, end_r + 1):
            for c_idx in range(openpyxl.utils.column_index_from_string(start_col), openpyxl.utils.column_index_from_string(end_col) + 1):
                c = ws_strat.cell(r_idx, c_idx)
                c.border = box_border
                if fill: c.fill = fill
                if font: c.font = font
                if alignment: c.alignment = alignment

    # ENCABEZADO
    ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
    cell = ws_strat[f'B{curr_row}']
    cell.value = "MONCHIS CAFÉ — ESTRATEGIA COMERCIAL DE VENTA Y JUSTIFICACIÓN FINANCIERA"
    cell.font = f_title
    cell.fill = fill_header
    cell.alignment = align_center
    apply_merged_borders('B', curr_row, 'G', curr_row, fill=fill_header)
    ws_strat.row_dimensions[curr_row].height = 26
    curr_row += 1
    
    ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
    cell = ws_strat[f'B{curr_row}']
    cell.value = "Actividad 4: Gestión de Finanzas | Proyecto Integrador de Desarrollo de Software"
    cell.font = f_subtitle
    cell.fill = fill_header
    cell.alignment = align_center
    apply_merged_borders('B', curr_row, 'G', curr_row, fill=fill_header)
    ws_strat.row_dimensions[curr_row].height = 18
    curr_row += 2
    
    # METADATOS
    meta = [
        ("Proyecto Integrador:", "Monchis Café — Plataforma Web y POS Seguro para Cafetería Orgánica", "Docente:", "Dr. Gabriel Navarro Salcedo"),
        ("Estudiante:", "Ernesto Hatuey Fierro Meléndez (Código: 215487963)", "Asignatura:", "Proyecto Integrador / Gestión de Finanzas"),
        ("Empresa Desarrolladora:", "Soluciones Digitales Fierro & Asociados S.A.S. de C.V.", "Fecha de Entrega:", "24 de Octubre de 2026")
    ]
    for lbl1, val1, lbl2, val2 in meta:
        ws_strat[f'B{curr_row}'].value = lbl1
        ws_strat[f'B{curr_row}'].font = f_bold
        ws_strat[f'B{curr_row}'].fill = fill_zebra
        ws_strat[f'B{curr_row}'].border = box_border
        
        ws_strat.merge_cells(f'C{curr_row}:D{curr_row}')
        ws_strat[f'C{curr_row}'].value = val1
        ws_strat[f'C{curr_row}'].font = f_regular
        apply_merged_borders('C', curr_row, 'D', curr_row)
        
        ws_strat[f'E{curr_row}'].value = lbl2
        ws_strat[f'E{curr_row}'].font = f_bold
        ws_strat[f'E{curr_row}'].fill = fill_zebra
        ws_strat[f'E{curr_row}'].border = box_border
        
        ws_strat.merge_cells(f'F{curr_row}:G{curr_row}')
        ws_strat[f'F{curr_row}'].value = val2
        ws_strat[f'F{curr_row}'].font = f_regular
        apply_merged_borders('F', curr_row, 'G', curr_row)
        
        ws_strat.row_dimensions[curr_row].height = 20
        curr_row += 1
        
    curr_row += 1

    # SECCIÓN 1: RESUMEN EJECUTIVO (FÓRMULAS VIVAS Y REFERENCIAS EXACTAS)
    ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
    sec1 = ws_strat[f'B{curr_row}']
    sec1.value = "1. RESUMEN EJECUTIVO DE COSTOS, PRECIO Y RESULTADOS FINANCIEROS"
    sec1.font = f_section
    sec1.fill = fill_section
    sec1.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    apply_merged_borders('B', curr_row, 'G', curr_row, fill=fill_section)
    ws_strat.row_dimensions[curr_row].height = 24
    curr_row += 1
    
    kpi_headers = [
        ("B", "Indicador Financiero Clave", align_left),
        ("C", "Valor (MXN)", align_center),
        ("D", "Referencia Celda", align_center),
        ("E", "% s/ Venta", align_center),
        ("F", "Interpretación Financiera", align_left),
        ("G", "Relevancia Operativa en el Proyecto", align_left)
    ]
    for col_let, h, al in kpi_headers:
        c = ws_strat[f'{col_let}{curr_row}']
        c.value = h
        c.font = f_tbl_hdr
        c.fill = fill_header
        c.alignment = al
        c.border = box_border
    ws_strat.row_dimensions[curr_row].height = 22
    curr_row += 1
    
    kpi_rows = [
        ("Costo Base Total (CTB)", "='RESUMEN Y COTIZACION'!B11", "'RESUMEN Y COTIZACION'!B11", "=C{r}/'RESUMEN Y COTIZACION'!$B$20", "Costo directo e indirecto acumulado.", "Nómina con 30% patronal, gastos fijos, nube, depreciación e intereses."),
        ("Fondo de Contingencia (10%)", "='RESUMEN Y COTIZACION'!B12", "'RESUMEN Y COTIZACION'!B12", "=C{r}/'RESUMEN Y COTIZACION'!$B$20", "Reserva para mitigación de riesgos.", "Amortigua variaciones cambiarias en servicios cloud y ajustes técnicos menores."),
        ("COSTO TOTAL DEL PROYECTO (CTP)", "='RESUMEN Y COTIZACION'!B13", "'RESUMEN Y COTIZACION'!B13", "=C{r}/'RESUMEN Y COTIZACION'!$B$20", "Inversión neta para producir la solución.", "Inversión indispensable para cubrir los 8 sprints de desarrollo y pruebas."),
        ("Margen de Utilidad sobre Ventas", "='RESUMEN Y COTIZACION'!B16", "'RESUMEN Y COTIZACION'!B16", "=C{r}", "Margen sobre precio fijado en 30.0%.", "Margen comercial sobre el precio facturado sin IVA, preservando rentabilidad."),
        ("PRECIO DE VENTA SIN IVA (COTIZACIÓN)", "='RESUMEN Y COTIZACION'!B17", "'RESUMEN Y COTIZACION'!B17", "=C{r}/'RESUMEN Y COTIZACION'!$B$20", "Precio comercial: CTP / (1 - Margen).", "Cotización base de la solución completa llave en mano con entrega en la nube."),
        ("IVA Trasladado (16%)", "='RESUMEN Y COTIZACION'!B18", "'RESUMEN Y COTIZACION'!B18", "=C{r}/'RESUMEN Y COTIZACION'!$B$20", "Impuesto fiscal acreditable por ley.", "Impuesto trasladado; acreditable para el cliente contra sus débitos de IVA."),
        ("PRECIO TOTAL FACTURADO CON IVA", "='RESUMEN Y COTIZACION'!B19", "'RESUMEN Y COTIZACION'!B19", "=C{r}/'RESUMEN Y COTIZACION'!$B$20", "Importe contractual total formalizado.", "Monto formal facturado suscrito en el contrato de prestación de servicios."),
        ("UTILIDAD ANTES DE ISR (PROYECTADA)", "='RESUMEN Y COTIZACION'!B24", "'RESUMEN Y COTIZACION'!B24", "=C{r}/'RESUMEN Y COTIZACION'!$B$20", "Ganancia de la empresa antes de ISR.", "Utilidad comercial antes del impuesto corporativo para reinversión y reservas."),
        ("ISR Corporativo Estimado (30%)", "=C{r_prev}*0.30", "C{r_prev} * 30%", "=C{r}/'RESUMEN Y COTIZACION'!$B$20", "Estimación de ISR según LISR (30%).", "Provisión del impuesto corporativo sobre la utilidad generada por el contrato."),
        ("UTILIDAD ESTIMADA DESPUÉS DE ISR", "=C{r_prev2}-C{r_prev}", "C{r_prev2} - C{r_prev}", "=C{r}/'RESUMEN Y COTIZACION'!$B$20", "Beneficio neto disponible (21.0%).", "Utilidad neta final de la empresa desarrolladora tras obligaciones fiscales."),
        ("Punto de Equilibrio (Soluciones)", "='RESUMEN Y COTIZACION'!B26", "'RESUMEN Y COTIZACION'!B26", "N/A", "1 proyecto único llave en mano.", "La comercialización del sistema al cliente cubre la estructura de costos."),
        ("Punto de Equilibrio en Pesos", "=ROUND(('RESUMEN Y COTIZACION'!B6+'RESUMEN Y COTIZACION'!B7+'RESUMEN Y COTIZACION'!B9+'RESUMEN Y COTIZACION'!B10)/(('RESUMEN Y COTIZACION'!B17-('RESUMEN Y COTIZACION'!B8+'RESUMEN Y COTIZACION'!B12))/'RESUMEN Y COTIZACION'!B17), 2)", "Fijos / Razón Margen Contribución", "=C{r}/'RESUMEN Y COTIZACION'!$B$17", "Facturación mínima de cobertura.", "Ingreso sin IVA necesario para cubrir la totalidad de costos fijos y variables."),
        ("Margen de Seguridad Comercial", "=ROUND(('RESUMEN Y COTIZACION'!B17-C{r_prev})/'RESUMEN Y COTIZACION'!B17, 4)", "(Precio - PE) / Precio", "=C{r}", "Holgura frente a variaciones (33.1%).", "Porcentaje en que puede reducirse la facturación antes de incurrir en pérdidas.")
    ]
    
    start_kpi_row = curr_row
    for label, f_val, ref, f_pct, interp, relev in kpi_rows:
        is_highlight = "PRECIO DE VENTA" in label or "UTILIDAD" in label or "COSTO TOTAL" in label
        row_fill = fill_success if "DESPUÉS DE ISR" in label or "ANTES DE ISR" in label else (fill_zebra if is_highlight else None)
        
        ws_strat[f'B{curr_row}'].value = label
        ws_strat[f'B{curr_row}'].font = f_bold if is_highlight else f_regular
        ws_strat[f'B{curr_row}'].border = box_border
        if row_fill: ws_strat[f'B{curr_row}'].fill = row_fill
        
        c_val = ws_strat[f'C{curr_row}']
        val_expr = f_val.replace("{r_prev2}", str(curr_row - 2)).replace("{r_prev}", str(curr_row - 1))
        c_val.value = val_expr
        c_val.font = f_bold if is_highlight else f_regular
        c_val.alignment = align_right
        c_val.border = box_border
        if "Margen" in label:
            c_val.number_format = fmt_pct
            c_val.alignment = align_center
        elif "Punto de Equilibrio (Soluciones)" in label:
            c_val.number_format = fmt_num
            c_val.alignment = align_center
        else:
            c_val.number_format = fmt_currency
        if row_fill: c_val.fill = row_fill
        
        c_ref = ws_strat[f'D{curr_row}']
        ref_expr = ref.replace("{r_prev2}", str(curr_row - 2)).replace("{r_prev}", str(curr_row - 1))
        c_ref.value = ref_expr
        c_ref.font = f_ref
        c_ref.alignment = align_center
        c_ref.border = box_border
        if row_fill: c_ref.fill = row_fill
        
        c_pct = ws_strat[f'E{curr_row}']
        pct_expr = f_pct.replace("{r}", str(curr_row)).replace("{r_prev}", str(curr_row - 1))
        c_pct.value = pct_expr
        c_pct.font = f_bold if is_highlight else f_regular
        c_pct.alignment = align_center
        c_pct.border = box_border
        if "{r}" in f_pct or "Margen" in label:
            c_pct.number_format = fmt_pct
        if row_fill: c_pct.fill = row_fill
        
        c_int = ws_strat[f'F{curr_row}']
        c_int.value = interp
        c_int.font = f_regular
        c_int.alignment = align_left
        c_int.border = box_border
        if row_fill: c_int.fill = row_fill
        
        c_rel = ws_strat[f'G{curr_row}']
        c_rel.value = relev
        c_rel.font = f_regular
        c_rel.alignment = align_left
        c_rel.border = box_border
        if row_fill: c_rel.fill = row_fill
        
        ws_strat.row_dimensions[curr_row].height = 20
        curr_row += 1
        
    curr_row += 1

    # SECCIÓN 2: CUADRO JUSTIFICATIVO DE COSTOS
    ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
    sec2 = ws_strat[f'B{curr_row}']
    sec2.value = "2. CUADRO JUSTIFICATIVO DE PERTINENCIA DE COSTOS DEL PROYECTO"
    sec2.font = f_section
    sec2.fill = fill_section
    sec2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    apply_merged_borders('B', curr_row, 'G', curr_row, fill=fill_section)
    ws_strat.row_dimensions[curr_row].height = 24
    curr_row += 1
    
    just_headers = [
        ("B", "Rubro Presupuestal", align_left),
        ("C", "Importe (MXN)", align_center),
        ("D", "% s/ CTP", align_center),
        ("E", "Clasificación", align_center),
        ("F", "Criterio Técnico y Justificación Operativa", align_left),
        ("G", "Impacto en Calidad y Mitigación de Riesgos", align_left)
    ]
    for col_let, h, al in just_headers:
        c = ws_strat[f'{col_let}{curr_row}']
        c.value = h
        c.font = f_tbl_hdr
        c.fill = fill_header
        c.alignment = al
        c.border = box_border
    ws_strat.row_dimensions[curr_row].height = 22
    curr_row += 1
    
    just_items = [
        ("Nómina y Carga Patronal (4 meses)",
         "='RESUMEN Y COTIZACION'!B6",
         "=C{r}/'RESUMEN Y COTIZACION'!$B$13",
         "Mano de Obra Directa",
         "4 perfiles de ingeniería: Líder/Scrum Master & DevOps ($36k), Frontend Vue 3 ($20k), Backend Node.js ($20k) y QA Ciberseguridad ($13.5k). Sueldos brutos de $89.5k/mes más estimación del 30% de carga patronal (IMSS, Infonavit 5%, ISN 3%).",
         "Permite la ejecución de 8 sprints bajo TDD (>80% cobertura) y reduce el riesgo de rotación técnica mediante compensaciones acordes al sector TI en la región."),
        
        ("Gastos Fijos Operativos Imputados",
         "='RESUMEN Y COTIZACION'!B7",
         "=C{r}/'RESUMEN Y COTIZACION'!$B$13",
         "Costo Indirecto Fijo",
         "Espacio coworking ($6k/mes), energía ($1.2k), fibra simétrica con IP fija ($1.5k), despacho contable ($2.5k), herramientas SaaS de colaboración ($1.8k), gestión comercial ($1k) y suministros ($1k), imputados al 100% en los 4 meses ($15k/mes).",
         "Provee infraestructura física, conectividad estable para CI/CD, soporte fiscal para timbrado CFDI 4.0 y herramientas de trabajo colaborativo."),
        
        ("Costos Directos del Proyecto (10 rubros)",
         "='RESUMEN Y COTIZACION'!B8",
         "=C{r}/'RESUMEN Y COTIZACION'!$B$13",
         "Costo Directo Variable",
         "GCP Cloud Run/SQL/Redis ($7.2k), GitHub/Sentry ($2.4k), CloudAMQP RabbitMQ ($2.2k), dominio/SSL ($950), dispositivos de prueba ($3k), pentesting externo ($4.5k), viáticos en La Piedad ($3k), papel térmico/cables ($850), SAST OWASP ($1.2k) y otros ($1.5k).",
         "Sustenta la arquitectura orientada a eventos con RabbitMQ y Patrón Saga, asegura el cumplimiento de seguridad en pagos y valida los periféricos en el mostrador físico."),
          
        ("Inversión Inicial Asignada (Depreciación)",
         "='RESUMEN Y COTIZACION'!B9",
         "=C{r}/'RESUMEN Y COTIZACION'!$B$13",
         "Depreciación de Activos",
         "Depreciación en 4 meses según vida útil fiscal: 4 laptops (36 meses), mobiliario (60 meses), red gigabit (36 meses), licencias IDEs (24 meses), capacitación RabbitMQ (12 meses) y kit POS de mostrador (24 meses) ($19,638.89 de $148.8k adquiridos).",
         "Aplica el principio contable de devengo: el contrato absorbe únicamente el desgaste proporcional de los bienes durante el tiempo de ejecución del proyecto."),
          
        ("Intereses por Financiamiento Puente",
         "='RESUMEN Y COTIZACION'!B10",
         "=C{r}/'RESUMEN Y COTIZACION'!$B$13",
         "Costo Financiero",
         "Crédito de capital de trabajo por $40,000.00 MXN al 18% anual a 24 meses. Se imputa el interés promedio lineal prorrateado a 4 meses ($1,321.19 MXN). Permite anticipar la adquisición de activos y absorber desfases en el cobro del anticipo.",
         "Evita retrasos operativos en el arranque derivados de demoras bancarias o administrativas en la recepción de la primera transferencia del cliente."),
          
        ("Fondo de Contingencia y Riesgos (10%)",
         "='RESUMEN Y COTIZACION'!B12",
         "=C{r}/'RESUMEN Y COTIZACION'!$B$13",
         "Gestión de Riesgos",
         "Reserva estratégica del 10% sobre el costo base ($57,316.01 MXN), asignada para mitigar fluctuaciones en tarifas de cómputo en la nube en USD, variaciones cambiarias y ajustes técnicos no previstos.",
         "Absorbe desviaciones operativas y constituye el único colchón financiero ante la ausencia de roles dedicados de DevOps, DBA o UX en el equipo."),
          
        ("TOTAL COSTO DEL PROYECTO (CTP)",
         "='RESUMEN Y COTIZACION'!B13",
         "=C{r}/'RESUMEN Y COTIZACION'!$B$13",
         "Costo Total Integral",
         "Inversión total requerida por la empresa desarrolladora para diseñar, programar, probar y desplegar en producción la plataforma web y POS Monchis Café.",
         "Sustenta una cotización estructurada y competitiva, respaldada en costos reales de ingeniería de software.")
    ]
    
    for rubro, f_monto, f_pct, tipo, crit, imp in just_items:
        is_tot = "TOTAL COSTO" in rubro
        row_fill = fill_subhead if is_tot else None
        
        ws_strat[f'B{curr_row}'].value = rubro
        ws_strat[f'B{curr_row}'].font = f_bold if is_tot else f_regular
        ws_strat[f'B{curr_row}'].border = box_border
        if row_fill: ws_strat[f'B{curr_row}'].fill = row_fill
        
        c_m = ws_strat[f'C{curr_row}']
        c_m.value = f_monto
        c_m.font = f_bold if is_tot else f_regular
        c_m.alignment = align_right
        c_m.number_format = fmt_currency
        c_m.border = box_border
        if row_fill: c_m.fill = row_fill
        
        c_p = ws_strat[f'D{curr_row}']
        c_p.value = f_pct.replace("{r}", str(curr_row))
        c_p.font = f_bold if is_tot else f_regular
        c_p.alignment = align_center
        c_p.number_format = fmt_pct
        c_p.border = box_border
        if row_fill: c_p.fill = row_fill
        
        c_t = ws_strat[f'E{curr_row}']
        c_t.value = tipo
        c_t.font = f_bold if is_tot else f_regular
        c_t.alignment = align_center
        c_t.border = box_border
        if row_fill: c_t.fill = row_fill
        
        c_c = ws_strat[f'F{curr_row}']
        c_c.value = crit
        c_c.font = f_regular
        c_c.alignment = align_left
        c_c.border = box_border
        if row_fill: c_c.fill = row_fill
        
        c_i = ws_strat[f'G{curr_row}']
        c_i.value = imp
        c_i.font = f_regular
        c_i.alignment = align_left
        c_i.border = box_border
        if row_fill: c_i.fill = row_fill
        
        ws_strat.row_dimensions[curr_row].height = 42 if not is_tot else 22
        curr_row += 1
        
    curr_row += 1

    # SECCIÓN 3: ESTRATEGIA DE VENTA COMERCIAL
    ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
    sec3 = ws_strat[f'B{curr_row}']
    sec3.value = "3. ESTRATEGIA DE VENTA COMERCIAL SEGÚN COSTO Y UTILIDAD"
    sec3.font = f_section
    sec3.fill = fill_section
    sec3.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    apply_merged_borders('B', curr_row, 'G', curr_row, fill=fill_section)
    ws_strat.row_dimensions[curr_row].height = 24
    curr_row += 1
    
    strat_text_blocks = [
        ("A. MODELO DE FIJACIÓN DE PRECIOS POR MARGEN SOBRE VENTAS",
         "Se determinó el precio mediante margen comercial sobre ventas antes de impuestos:\n"
         "• Fórmula de Fijación: Precio de Venta = CTP / (1 - Margen) = $630,476.09 / (1 - 0.30) = $900,680.12 MXN sin IVA ($1,044,788.94 MXN con IVA).\n"
         "• Margen vs Recargo (Markup): La fórmula de margen asegura que la utilidad antes de ISR ($270,204.04 MXN) represente con exactitud el 30.0% del precio comercial facturado antes de impuestos, a diferencia de un recargo simple que reduce la rentabilidad al 23.1%.\n"
         "• Posicionamiento Comercial: Se posiciona como una solución de software a medida, desacoplada mediante eventos RabbitMQ y trazabilidad ODS 12, superando las limitaciones de los sistemas POS genéricos de caja."),

        ("B. MODELO DE MONETIZACIÓN HÍBRIDO (IMPLEMENTACIÓN + INGRESOS RECURRENTES)",
         "1. Venta Inicial Llave en Mano ($900,680.12 MXN + IVA): Licenciamiento perpetuo de la obra de software sobre infraestructura propia del cliente en Google Cloud Platform, configuración de microservicios, migración de catálogos y capacitación presencial a baristas y cajeros.\n"
         "2. Póliza de Soporte, Mantenimiento Evolutivo y SLA ($18,500.00 MXN + IVA/mes a partir del Mes 5):\n"
         "    - Disponibilidad del 99.5% con monitoreo de microservicios y Dead Letter Queue (DLQ).\n"
         "    - Ventana de respuesta a incidentes críticos menor a 2 horas durante horario de servicio de la cafetería.\n"
         "    - Actualizaciones menores de dependencias, parches de seguridad y optimizaciones SEO continuas.\n"
         "    - Margen operativo de la póliza: 55%, generando un ingreso recurrente anual (ARR) de $222,000.00 MXN.\n"
         "3. Módulo de Franquicias / Nuevas Sucursales: Cuota de licenciamiento de $45,000.00 MXN por nueva tienda más $3,500.00 MXN/mes de soporte y mantenimiento por sucursal, escalando con el crecimiento del cliente."),

        ("C. PROPUESTA DE VALOR Y RETORNO DE INVERSIÓN (PAYBACK) PARA MONCHIS CAFÉ",
         "La inversión se sustenta en beneficios operativos mensuales proyectados para la cafetería:\n"
         "• Reducción de Mermas de Café Orgánico (-22%): Ahorro mensual estimado de $14,500.00 MXN mediante rotación PEPS y control de lotes.\n"
         "• Incremento en Ticket Promedio (+24%): Ingreso mensual adicional de $38,500.00 MXN por sellos digitales Monchis Rewards y bono termo.\n"
         "• Eliminación de Fugas en Caja (-98%): Recuperación de $9,200.00 MXN mensuales por conciliación automática de pagos mixtos y control de cambio.\n"
         "• Captura de Nuevos Clientes por SEO Local: Ingreso mensual de $16,800.00 MXN derivado de tráfico de Google Maps e Instagram.\n"
         "• Total Beneficio Operativo Bruto Estimado: $79,000.00 MXN/mes ($948,000.00 MXN anuales).\n"
         "Nota de supuestos: Las cifras anteriores son estimaciones preliminares sujetas a validación con datos operativos históricos reales del negocio.\n"
         "Evaluación del Período de Recuperación (Payback Period):\n"
         "1. Payback Bruto: $900,680.12 / $79,000.00 = 11.4 meses (sin descontar costo de póliza ni costo de ventas).\n"
         "2. Payback Neto con Póliza SLA: $900,680.12 / ($79,000.00 - $18,500.00) = $900,680.12 / $60,500.00 = 14.9 meses.\n"
         "3. Payback Prudente (60% margen de contribución en ventas adicionales): Beneficio neto de $56,880.00/mes ($33,180 margen ventas + $23,700 ahorros). Plazo sin póliza: 15.8 meses; plazo neto con póliza ($38,380/mes): 23.5 meses."),

        ("D. ESQUEMA DE CIERRE Y FINANCIAMIENTO POR HITOS ÁGILES (6 HITOS DE ENTREGA)",
         "Para dar certidumbre al cliente y asegurar flujo de efectivo para el equipo, el cobro se divide en 6 hitos contra entregables en Staging:\n"
         "• Hito 0 (Firma y Arranque): Anticipo del 20% ($180,136.02 MXN sin IVA). Setup de monorepo, Docker Compose, esquema Prisma y licencias.\n"
         "• Hito 1 (Sprint 2 - Backend & Seguridad): 15% ($135,102.02 MXN sin IVA). Backend Core, JWT stateless, RBAC y reCAPTCHA en Staging.\n"
         "• Hito 2 (Sprint 4 - Eventos & Frontend): 15% ($135,102.02 MXN sin IVA). Broker RabbitMQ CloudAMQP, Patrón Saga y vistas Vue 3 prerenderizadas.\n"
         "• Hito 3 (Sprint 6 - POS & Trazabilidad): 15% ($135,102.02 MXN sin IVA). POS táctil de barra, pagos mixtos, escaneo QR y trazabilidad ODS 12.\n"
         "• Hito 4 (Sprint 8 - Analítica & Testing): 15% ($135,102.02 MXN sin IVA). Dashboard de ventas con UTM, suites Vitest/Playwright y reporte SAST.\n"
         "• Hito 5 (Cierre y Despliegue en GCP): Finiquito del 20% ($180,136.02 MXN sin IVA). Terraform en GCP, manuales de usuario y entrega de llaves.\n"
         "Nota sobre financiamiento puente: La línea de crédito de $40,000.00 MXN se justifica para anticipar la compra de activos de laboratorio ($148.8k) y depósitos cloud antes de la dispersión bancaria del anticipo, protegiendo el calendario de trabajo frente a demoras administrativas.")
    ]
    
    for sub_title, text_content in strat_text_blocks:
        ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
        st_cell = ws_strat[f'B{curr_row}']
        st_cell.value = sub_title
        st_cell.font = f_subhead
        st_cell.fill = fill_subhead
        apply_merged_borders('B', curr_row, 'G', curr_row, fill=fill_subhead)
        ws_strat.row_dimensions[curr_row].height = 22
        curr_row += 1
        
        ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
        tb_cell = ws_strat[f'B{curr_row}']
        tb_cell.value = text_content
        tb_cell.font = f_regular
        tb_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        apply_merged_borders('B', curr_row, 'G', curr_row)
        
        line_count = text_content.count('\n') + 1
        ws_strat.row_dimensions[curr_row].height = max(58, line_count * 17)
        curr_row += 2

    # SECCIÓN 4: ANÁLISIS DE SENSIBILIDAD Y ESCENARIOS
    ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
    sec4 = ws_strat[f'B{curr_row}']
    sec4.value = "4. ANÁLISIS DE SENSIBILIDAD Y ESCENARIOS DE RENTABILIDAD"
    sec4.font = f_section
    sec4.fill = fill_section
    sec4.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    apply_merged_borders('B', curr_row, 'G', curr_row, fill=fill_section)
    ws_strat.row_dimensions[curr_row].height = 24
    curr_row += 1
    
    sens_headers = [
        ("B", "Escenario Comercial", align_left),
        ("C", "Margen s/ Venta", align_center),
        ("D", "Precio sin IVA (MXN)", align_center),
        ("E", "Precio con IVA", align_center),
        ("F", "Utilidad antes ISR", align_center),
        ("G", "Criterio y Táctica de Negociación", align_left)
    ]
    for col_let, h, al in sens_headers:
        c = ws_strat[f'{col_let}{curr_row}']
        c.value = h
        c.font = f_tbl_hdr
        c.fill = fill_header
        c.alignment = al
        c.border = box_border
    ws_strat.row_dimensions[curr_row].height = 22
    curr_row += 1
    
    sens_rows = [
        ("Escenario Conservador (Piso Mínimo)",
         0.20,
         "=ROUND('RESUMEN Y COTIZACION'!$B$13/(1-C{r}), 2)",
         "=ROUND(D{r}*(1+PARAMETROS!$B$7), 2)",
         "=ROUND(D{r}-'RESUMEN Y COTIZACION'!$B$13, 2)",
         "Piso de negociación que sacrifica margen comercial (del 30% al 20%), no reducción de costos de desarrollo. Asegura el contrato ante restricciones del cliente."),
        
        ("Escenario Esperado / Base (Recomendado)",
         "=PARAMETROS!$B$6",
         "='RESUMEN Y COTIZACION'!$B$17",
         "='RESUMEN Y COTIZACION'!$B$19",
         "='RESUMEN Y COTIZACION'!$B$24",
         "Propuesta recomendada. Mantiene el margen del 30.0% proyectado para la viabilidad de la empresa y amortización equilibrada para el cliente."),
        
        ("Escenario Premium / Valor Agregado",
         0.35,
         "=ROUND('RESUMEN Y COTIZACION'!$B$13/(1-C{r}), 2)",
         "=ROUND(D{r}*(1+PARAMETROS!$B$7), 2)",
         "=ROUND(D{r}-'RESUMEN Y COTIZACION'!$B$13, 2)",
         "Paquete con 6 meses de póliza SLA 99.5% ($49,950 MXN costo asumido) y analítica avanzada. Utilidad ajustada post-SLA: $289,537.12 MXN (margen efectivo del 29.9%).")
    ]
    
    for esc, mrg, f_psi, f_pci, f_ut, crit in sens_rows:
        is_base = "Esperado" in esc
        row_fill = fill_zebra if is_base else None
        
        ws_strat[f'B{curr_row}'].value = esc
        ws_strat[f'B{curr_row}'].font = f_bold if is_base else f_regular
        ws_strat[f'B{curr_row}'].border = box_border
        if row_fill: ws_strat[f'B{curr_row}'].fill = row_fill
        
        c_m = ws_strat[f'C{curr_row}']
        c_m.value = mrg
        c_m.font = f_bold if is_base else f_regular
        c_m.alignment = align_center
        c_m.number_format = fmt_pct
        c_m.border = box_border
        if row_fill: c_m.fill = row_fill
        
        c_psi = ws_strat[f'D{curr_row}']
        c_psi.value = f_psi.replace("{r}", str(curr_row)) if "{r}" in str(f_psi) else f_psi
        c_psi.font = f_bold if is_base else f_regular
        c_psi.alignment = align_right
        c_psi.number_format = fmt_currency
        c_psi.border = box_border
        if row_fill: c_psi.fill = row_fill
        
        c_pci = ws_strat[f'E{curr_row}']
        c_pci.value = f_pci.replace("{r}", str(curr_row)) if "{r}" in str(f_pci) else f_pci
        c_pci.font = f_bold if is_base else f_regular
        c_pci.alignment = align_right
        c_pci.number_format = fmt_currency
        c_pci.border = box_border
        if row_fill: c_pci.fill = row_fill
        
        c_ut = ws_strat[f'F{curr_row}']
        c_ut.value = f_ut.replace("{r}", str(curr_row)) if "{r}" in str(f_ut) else f_ut
        c_ut.font = f_bold if is_base else f_regular
        c_ut.alignment = align_right
        c_ut.number_format = fmt_currency
        c_ut.border = box_border
        if row_fill: c_ut.fill = row_fill
        
        c_cr = ws_strat[f'G{curr_row}']
        c_cr.value = crit
        c_cr.font = f_regular
        c_cr.alignment = align_left
        c_cr.border = box_border
        if row_fill: c_cr.fill = row_fill
        
        ws_strat.row_dimensions[curr_row].height = 28
        curr_row += 1
        
    curr_row += 1

    # SECCIÓN 5: NOTAS TÉCNICAS DE RIESGO Y SENSIBILIDAD FISCAL
    ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
    sec5 = ws_strat[f'B{curr_row}']
    sec5.value = "5. NOTAS TÉCNICAS DE RIESGO OPERATIVO Y SENSIBILIDAD FISCAL"
    sec5.font = f_section
    sec5.fill = fill_section
    sec5.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    apply_merged_borders('B', curr_row, 'G', curr_row, fill=fill_section)
    ws_strat.row_dimensions[curr_row].height = 24
    curr_row += 1
    
    risk_notes = (
        "• Nota de Riesgo de Perfiles Técnicos: El proyecto no dispone de perfiles dedicados y exclusivos de DevOps, DBA ni diseñador UX; dichas responsabilidades son absorbidas por los cuatro ingenieros de desarrollo. Por tanto, el fondo de contingencia del 10% ($57,316.01 MXN) constituye el único margen disponible para absorber desviaciones técnicas o retrabajos en infraestructura.\n"
        "• Nota de Sensibilidad de Carga Patronal: Por cada 5.0 puntos porcentuales adicionales de carga patronal sobre la nómina ($17,900.00 MXN en sueldos brutos + $1,790.00 MXN de contingencia = $19,690.00 MXN en CTP), el precio de venta sin IVA se incrementa en $28,128.57 MXN (≈ $28,100 MXN sin IVA) para conservar el margen del 30.0%.\n"
        "• Tratamiento Tributario del IVA: El IVA facturado ($144,108.82 MXN) corresponde a un impuesto fiscal trasladado que no forma parte del ingreso ni de la utilidad de la empresa desarrolladora, siendo acreditable para el cliente contra sus débitos fiscales."
    )
    ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
    rn_cell = ws_strat[f'B{curr_row}']
    rn_cell.value = risk_notes
    rn_cell.font = f_regular
    rn_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    apply_merged_borders('B', curr_row, 'G', curr_row)
    ws_strat.row_dimensions[curr_row].height = 54
    curr_row += 2

    # Pie institucional sobrio
    ws_strat.merge_cells(f'B{curr_row}:G{curr_row}')
    foot = ws_strat[f'B{curr_row}']
    foot.value = "Actividad 4: Gestión de Finanzas | Cátedra: Proyecto Integrador de Desarrollo de Software | Docente: Dr. Gabriel Navarro Salcedo | Centro Universitario de Tonalá (CUTonalá) — Universidad de Guadalajara."
    foot.font = f_italic
    foot.alignment = align_center
    apply_merged_borders('B', curr_row, 'G', curr_row)

    # -------------------------------------------------------------
    # 8. NUEVA HOJA: FLUJO DE CAJA MENSUAL SIMPLE
    # -------------------------------------------------------------
    cf_sheet_name = 'FLUJO DE CAJA'
    if cf_sheet_name in wb.sheetnames:
        del wb[cf_sheet_name]
    
    ws_cf = wb.create_sheet(title=cf_sheet_name)
    ws_cf.views.sheetView[0].showGridLines = True
    
    cf_widths = {'A': 4, 'B': 36, 'C': 18, 'D': 18, 'E': 18, 'F': 18, 'G': 20, 'H': 4}
    for col, width in cf_widths.items():
        ws_cf.column_dimensions[col].width = width
        
    cf_row = 2
    
    # Banner Flujo de Caja
    ws_cf.merge_cells(f'B{cf_row}:G{cf_row}')
    c_title = ws_cf[f'B{cf_row}']
    c_title.value = "MONCHIS CAFÉ — PROYECCIÓN DE FLUJO DE CAJA OPERATIVO (4 MESES)"
    c_title.font = f_title
    c_title.fill = fill_header
    c_title.alignment = align_center
    for col_idx in range(2, 8):
        ws_cf.cell(cf_row, col_idx).border = box_border
        ws_cf.cell(cf_row, col_idx).fill = fill_header
    ws_cf.row_dimensions[cf_row].height = 26
    cf_row += 1
    
    ws_cf.merge_cells(f'B{cf_row}:G{cf_row}')
    c_sub = ws_cf[f'B{cf_row}']
    c_sub.value = "Control Mensual de Entradas (Cobros con IVA) vs Salidas Operativas y Amortización"
    c_sub.font = f_subtitle
    c_sub.fill = fill_header
    c_sub.alignment = align_center
    for col_idx in range(2, 8):
        ws_cf.cell(cf_row, col_idx).border = box_border
        ws_cf.cell(cf_row, col_idx).fill = fill_header
    ws_cf.row_dimensions[cf_row].height = 18
    cf_row += 2
    
    # Headers de la tabla
    cf_headers = ["Concepto de Flujo de Efectivo", "Mes 1 (MXN)", "Mes 2 (MXN)", "Mes 3 (MXN)", "Mes 4 (MXN)", "Total Proyecto"]
    for col_idx, h_text in enumerate(cf_headers, start=2):
        c = ws_cf.cell(cf_row, col_idx, value=h_text)
        c.font = f_tbl_hdr
        c.fill = fill_header
        c.alignment = align_center if col_idx > 2 else align_left
        c.border = box_border
    ws_cf.row_dimensions[cf_row].height = 22
    cf_row += 1
    
    cf_data = [
        # ENTRADAS
        ("1. INGRESOS POR HITOS DE FACTURACIÓN (CON IVA)", None, None, None, None, None, "SECTION"),
        ("Cobro de Hito 0: Anticipo Inicial (20%)", "=ROUND('RESUMEN Y COTIZACION'!$B$19*0.20, 2)", 0, 0, 0, "=SUM(C{r}:F{r})", "ROW"),
        ("Cobro de Hito 1: Sprint 2 Backend Core (15%)", "=ROUND('RESUMEN Y COTIZACION'!$B$19*0.15, 2)", 0, 0, 0, "=SUM(C{r}:F{r})", "ROW"),
        ("Cobro de Hito 2: Sprint 4 RabbitMQ & Vue 3 (15%)", 0, "=ROUND('RESUMEN Y COTIZACION'!$B$19*0.15, 2)", 0, 0, "=SUM(C{r}:F{r})", "ROW"),
        ("Cobro de Hito 3: Sprint 6 POS Mostrador (15%)", 0, 0, "=ROUND('RESUMEN Y COTIZACION'!$B$19*0.15, 2)", 0, "=SUM(C{r}:F{r})", "ROW"),
        ("Cobro de Hito 4: Sprint 8 Analítica & QA (15%)", 0, 0, 0, "=ROUND('RESUMEN Y COTIZACION'!$B$19*0.15, 2)", "=SUM(C{r}:F{r})", "ROW"),
        ("Cobro de Hito 5: Finiquito y Cierre en GCP (20%)", 0, 0, 0, "=ROUND('RESUMEN Y COTIZACION'!$B$19*0.20, 2)", "=SUM(C{r}:F{r})", "ROW"),
        ("Financiamiento Puente Inicial (Crédito)", "=PRESTAMO!$B$2", 0, 0, 0, "=SUM(C{r}:F{r})", "ROW"),
        ("TOTAL ENTRADAS DE EFECTIVO (A)", "=SUM(C{r_start}:C{r_prev})", "=SUM(D{r_start}:D{r_prev})", "=SUM(E{r_start}:E{r_prev})", "=SUM(F{r_start}:F{r_prev})", "=SUM(G{r_start}:G{r_prev})", "SUBTOTAL"),
        
        # SALIDAS
        ("2. EGRESOS OPERATIVOS DEL PROYECTO", None, None, None, None, None, "SECTION"),
        ("Nómina y Cargas Patronales (30% IMSS/INFONAVIT/ISN)", "=ROUND('RESUMEN Y COTIZACION'!$B$6/4, 2)", "=ROUND('RESUMEN Y COTIZACION'!$B$6/4, 2)", "=ROUND('RESUMEN Y COTIZACION'!$B$6/4, 2)", "=ROUND('RESUMEN Y COTIZACION'!$B$6/4, 2)", "=SUM(C{r}:F{r})", "ROW"),
        ("Gastos Fijos de Oficina y Servicios", "=ROUND('RESUMEN Y COTIZACION'!$B$7/4, 2)", "=ROUND('RESUMEN Y COTIZACION'!$B$7/4, 2)", "=ROUND('RESUMEN Y COTIZACION'!$B$7/4, 2)", "=ROUND('RESUMEN Y COTIZACION'!$B$7/4, 2)", "=SUM(C{r}:F{r})", "ROW"),
        ("Costos Directos de Nube, Licencias y Seguridad", "=ROUND('RESUMEN Y COTIZACION'!$B$8/4, 2)", "=ROUND('RESUMEN Y COTIZACION'!$B$8/4, 2)", "=ROUND('RESUMEN Y COTIZACION'!$B$8/4, 2)", "=ROUND('RESUMEN Y COTIZACION'!$B$8/4, 2)", "=SUM(C{r}:F{r})", "ROW"),
        ("Adquisición Inicial de Activos de Laboratorio", "='INVERSION INICIAL'!F8", 0, 0, 0, "=SUM(C{r}:F{r})", "ROW"),
        ("Amortización Mensual del Préstamo Bancario", "=ROUND('PRESTAMO'!$B$4, 2)", "=ROUND('PRESTAMO'!$B$4, 2)", "=ROUND('PRESTAMO'!$B$4, 2)", "=ROUND('PRESTAMO'!$B$4, 2)", "=SUM(C{r}:F{r})", "ROW"),
        ("TOTAL SALIDAS DE EFECTIVO (B)", "=SUM(C{r_start}:C{r_prev})", "=SUM(D{r_start}:D{r_prev})", "=SUM(E{r_start}:E{r_prev})", "=SUM(F{r_start}:F{r_prev})", "=SUM(G{r_start}:G{r_prev})", "SUBTOTAL"),
        
        # RESULTADO NETO
        ("3. FLUJO NETO Y POSICIÓN DE CAJA", None, None, None, None, None, "SECTION"),
        ("FLUJO NETO MENSUAL (A - B)", "=C{r_entradas}-C{r_salidas}", "=D{r_entradas}-D{r_salidas}", "=E{r_entradas}-E{r_salidas}", "=F{r_entradas}-F{r_salidas}", "=G{r_entradas}-G{r_salidas}", "TOTAL_NET"),
        ("SALDO DE CAJA ACUMULADO AL CIERRE", "=C{r_prev}", "=C{r}+D{r_prev}", "=D{r}+E{r_prev}", "=E{r}+F{r_prev}", "=F{r}", "FINAL_BALANCE")
    ]
    
    r_in_start = 0
    r_in_end = 0
    r_out_start = 0
    r_out_end = 0
    row_total_in = 0
    row_total_out = 0
    row_net = 0
    
    for item in cf_data:
        label, m1, m2, m3, m4, tot, row_type = item
        
        if row_type == "SECTION":
            ws_cf.merge_cells(f'B{cf_row}:G{cf_row}')
            c = ws_cf[f'B{cf_row}']
            c.value = label
            c.font = f_subhead
            c.fill = fill_subhead
            for col_idx in range(2, 8):
                ws_cf.cell(cf_row, col_idx).border = box_border
                ws_cf.cell(cf_row, col_idx).fill = fill_subhead
            ws_cf.row_dimensions[cf_row].height = 20
            if "1. INGRESOS" in label:
                r_in_start = cf_row + 1
            elif "2. EGRESOS" in label:
                r_out_start = cf_row + 1
            cf_row += 1
            continue
            
        is_sub = row_type == "SUBTOTAL"
        is_net = row_type in ["TOTAL_NET", "FINAL_BALANCE"]
        fill_curr = fill_zebra if is_sub else (fill_success if is_net else None)
        font_curr = f_bold if (is_sub or is_net) else f_regular
        
        ws_cf[f'B{cf_row}'].value = label
        ws_cf[f'B{cf_row}'].font = font_curr
        ws_cf[f'B{cf_row}'].border = box_border
        if fill_curr: ws_cf[f'B{cf_row}'].fill = fill_curr
        
        vals = [m1, m2, m3, m4, tot]
        for col_idx, v in enumerate(vals, start=3):
            col_letter = get_column_letter(col_idx)
            cell = ws_cf[f'{col_letter}{cf_row}']
            
            if isinstance(v, str):
                expr = v.replace("{r}", str(cf_row))\
                        .replace("{r_start}", str(r_in_start if "TOTAL ENTRADAS" in label else r_out_start))\
                        .replace("{r_prev}", str(cf_row - 1))\
                        .replace("{r_entradas}", str(row_total_in))\
                        .replace("{r_salidas}", str(row_total_out))
                cell.value = expr
            else:
                cell.value = v
                
            cell.font = font_curr
            cell.alignment = align_right
            cell.number_format = fmt_currency
            cell.border = box_border
            if fill_curr: cell.fill = fill_curr
            
        ws_cf.row_dimensions[cf_row].height = 20
        
        if label == "TOTAL ENTRADAS DE EFECTIVO (A)":
            row_total_in = cf_row
        elif label == "TOTAL SALIDAS DE EFECTIVO (B)":
            row_total_out = cf_row
        elif label == "FLUJO NETO MENSUAL (A - B)":
            row_net = cf_row
            
        cf_row += 1

    # -------------------------------------------------------------
    # 9. GUARDAR ARCHIVO EN TODAS LAS UBICACIONES CORRESPONDIENTES
    # -------------------------------------------------------------
    downloads_dir = r"C:\Users\erfierro\Downloads"
    proj_presupuesto_dir = r"d:\ernestofm\Proyecto\docs\presupuesto"
    proj_entregables_dir = r"d:\ernestofm\Proyecto\docs\entregables"
    os.makedirs(proj_presupuesto_dir, exist_ok=True)
    os.makedirs(proj_entregables_dir, exist_ok=True)
    
    paths_to_save = [
        os.path.join(downloads_dir, "Actividad_4_Gestion_Finanzas_MonchisCafe.xlsx"),
        os.path.join(downloads_dir, "ERNESTO HATUEY FIERRO MELENDEZ - ACTIVIDAD 4 GESTION DE FINANZAS.xlsx"),
        os.path.join(proj_entregables_dir, "Actividad_4_Gestion_Finanzas_MonchisCafe.xlsx"),
        os.path.join(proj_presupuesto_dir, "Actividad_4_Gestion_Finanzas_MonchisCafe.xlsx"),
        os.path.join(proj_presupuesto_dir, "ERNESTO HATUEY FIERRO MELENDEZ - GESTION DE FINANZAS (COMPLETO).xlsx")
    ]
    
    for p in paths_to_save:
        wb.save(p)
        print(f"[OK] Archivo guardado exitosamente en: {p}")

if __name__ == "__main__":
    build_financial_model()
