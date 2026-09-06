import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def populate_smartsheet_template(source_template_path, output_project_path, output_downloads_path):
    # Cargar la plantilla original de Smartsheet
    wb = openpyxl.load_workbook(source_template_path)
    ws = wb['Project Budget Proposal']

    # 1. Título principal
    ws['B1'].value = "MONCHIS CAFÉ — SOFTWARE PROJECT BUDGET PROPOSAL (ALINEADO A ODS)"

    # 2. BLOQUE 1 (PROJECT 1): DESARROLLO DE SOFTWARE Y SPRINTS
    ws['B4'].value = "P1"
    ws['C4'].value = "DESARROLLO DE SOFTWARE Y SPRINTS (MANO DE OBRA)"

    p1_tasks = [
        # (fila, wbs, task, desc, status, p_start, a_start, end, hr, rate, budget)
        (5, "1.0", "Arquitectura Monorepo y Modelo BD", "Setup Turborepo, Docker Compose y esquema Prisma PostgreSQL", "Completed", "2026-08-31", "2026-08-31", "2026-09-14", 160, 150.00, 24000.00),
        (6, "2.0", "Backend Core y Seguridad Stateless", "JWT rotativo en cookies seguras, RBAC y Google reCAPTCHA v2/v3", "Completed", "2026-09-15", "2026-09-15", "2026-09-28", 240, 125.00, 30000.00),
        (7, "3.0", "Broker RabbitMQ y Patrón Saga", "Topic Exchange cafeteria.events, orquestador reversas y DLQ", "Completed", "2026-09-29", "2026-09-29", "2026-10-12", 240, 125.00, 30000.00),
        (8, "3.1", "Frontend Core y Prerenderizado SEO", "Vistas Vue 3, tokens pastel, vite-ssg y Schema.org LocalBusiness", "Completed", "2026-10-13", "2026-10-13", "2026-10-26", 240, 125.00, 30000.00),
        (9, "3.2", "Punto de Venta (POS) y Pagos Mixtos", "Interfaz táctil barra, escaneo QR, cobro efectivo/tarjeta/SPEI y vuelto", "In Progress", "2026-10-27", "2026-10-27", "2026-11-09", 400, 125.00, 50000.00),
        (10, "3.3", "Trazabilidad ODS y Monchis Rewards", "Catálogo fincas café regional, bono termo -$5 y sellos digitales", "In Progress", "2026-11-10", "2026-11-10", "2026-11-23", 320, 125.00, 40000.00),
        (11, "3.4", "Dashboard Admin y Atribución UTM", "Panel analítico de ventas por canal Google Maps vs Instagram y DLQ", "Planned", "2026-11-24", "2026-11-24", "2026-12-07", 240, 125.00, 30000.00),
        (12, "4.0", "Testing Multicapa y Ciberseguridad", "Suites Vitest, Playwright 3 viewports, escaneo OWASP y a11y WCAG", "Planned", "2026-11-24", "2026-11-24", "2026-12-07", 480, 112.50, 54000.00),
        (13, "5.0", "Despliegue Cloud en GCP y Release", "Infraestructura serverless con Terraform (Cloud Run, SQL) y release", "Planned", "2026-12-08", "2026-12-08", "2026-12-21", 400, 175.00, 70000.00),
    ]

    for r, wbs, task, desc, status, ps, as_, ed, hr, rate, budget in p1_tasks:
        ws[f'B{r}'].value = wbs
        ws[f'C{r}'].value = task
        ws[f'D{r}'].value = desc
        ws[f'E{r}'].value = status
        ws[f'F{r}'].value = ps
        ws[f'G{r}'].value = as_
        ws[f'H{r}'].value = ed
        ws[f'I{r}'].value = hr
        ws[f'J{r}'].value = rate
        ws[f'P{r}'].value = budget
        # Las celdas Q y R ya tienen sus fórmulas nativas =((I*J)+... ) y =Q-P

    # 3. BLOQUE 2 (PROJECT 2): HARDWARE, NUBE Y CONTINGENCIA
    ws['B16'].value = "P2"
    ws['C16'].value = "EQUIPAMIENTO, INFRAESTRUCTURA CLOUD Y CONTINGENCIA"

    p2_tasks = [
        # (fila, wbs, task, desc, status, ps, as_, ed, units, unit_cost, travel, equip_space, misc, budget)
        (17, "1", "Terminal Táctil Mostrador POS", "Tablet Android 10.5\" con soporte metálico de mostrador", "Purchased", "2026-08-31", "2026-08-31", "2026-09-07", 1, 6500.00, 0, 0, 0, 6500.00),
        (18, "2", "Impresora Térmica de Tickets", "Impresora 80mm USB/Red para tickets de papel reciclado", "Purchased", "2026-08-31", "2026-08-31", "2026-09-07", 1, 1800.00, 0, 0, 0, 1800.00),
        (19, "3", "Lector de Código de Barras / QR", "Scanner 1D/2D USB para inventario y sellos de lealtad", "Purchased", "2026-08-31", "2026-08-31", "2026-09-07", 1, 1200.00, 0, 0, 0, 1200.00),
        (20, "4", "Google Cloud Platform (GCP)", "Cloud Run, Cloud SQL PostgreSQL y Memorystore Redis (4 meses)", "Active", "2026-08-31", "2026-08-31", "2026-12-21", 0, 0, 0, 7200.00, 0, 7200.00),
        (21, "5", "CloudAMQP (RabbitMQ Administrado)", "Broker de colas para Patrón Saga con retardo y DLQ (4 meses)", "Active", "2026-08-31", "2026-08-31", "2026-12-21", 0, 0, 0, 2200.00, 0, 2200.00),
        (22, "6", "Dominio Web y Certificados SSL", "Registro de dominio oficial y certificados TLS 1.3 (1 año)", "Active", "2026-08-31", "2026-08-31", "2026-12-21", 0, 0, 0, 0, 950.00, 950.00),
        (23, "7", "Licencias Repositorio y DevOps", "GitHub Team y licencias de herramientas colaborativas (4 meses)", "Active", "2026-08-31", "2026-08-31", "2026-12-21", 0, 0, 0, 0, 2400.00, 2400.00),
        (24, "8", "Fondo de Contingencia y Riesgos", "Reserva del 10% para imprevistos técnicos y picos de tráfico", "Allocated", "2026-08-31", "2026-08-31", "2026-12-21", 0, 0, 0, 0, 38025.00, 38025.00),
        (25, "9", "Capacitación Operativa en Barra", "Manuales y entrenamiento al personal barista en POS táctil", "Planned", "2026-12-14", "2026-12-14", "2026-12-21", 0, 0, 0, 0, 0.00, 0.00),
    ]

    for r, wbs, task, desc, status, ps, as_, ed, units, u_cost, trav, eq_sp, misc, budget in p2_tasks:
        ws[f'B{r}'].value = wbs
        ws[f'C{r}'].value = task
        ws[f'D{r}'].value = desc
        ws[f'E{r}'].value = status
        ws[f'F{r}'].value = ps
        ws[f'G{r}'].value = as_
        ws[f'H{r}'].value = ed
        if units > 0:
            ws[f'K{r}'].value = units
            ws[f'L{r}'].value = u_cost
        if trav > 0:
            ws[f'M{r}'].value = trav
        if eq_sp > 0:
            ws[f'N{r}'].value = eq_sp
        if misc > 0:
            ws[f'O{r}'].value = misc
        ws[f'P{r}'].value = budget

    # 4. Fila 28: TOTAL GENERAL DEL PROYECTO MONCHIS CAFÉ
    total_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    total_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    
    ws['B28'].value = "TOTAL"
    ws['C28'].value = "PRESUPUESTO TOTAL INTEGRAL MONCHIS CAFÉ (P1 + P2)"
    ws['H28'].value = "GRAN TOTAL"
    ws['P28'].value = "=P14+P26"
    ws['Q28'].value = "=Q14+Q26"
    ws['R28'].value = "=R14+R26"

    for col in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R']:
        cell = ws[f'{col}28']
        cell.fill = total_fill
        cell.font = total_font
        if col in ['P', 'Q', 'R']:
            cell.number_format = '$#,##0.00'

    # Guardar en las rutas
    wb.save(output_project_path)
    wb.save(output_downloads_path)
    print(f"Plantilla oficial completada exitosamente:")
    print(f"  -> En proyecto: {output_project_path}")
    print(f"  -> En descargas: {output_downloads_path}")

if __name__ == "__main__":
    src = r"C:\Users\erfierro\Downloads\IC-Project-Budget-Proposal-11292.xlsx"
    out_proj = r"d:\ernestofm\Proyecto\docs\presupuesto\Presupuesto_Monchis_Cafe.xlsx"
    out_down = r"C:\Users\erfierro\Downloads\IC-Project-Budget-Proposal-11292.xlsx"
    populate_smartsheet_template(src, out_proj, out_down)
